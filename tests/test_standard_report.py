from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import yaml
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"


def run_cli(script: Path, *args: object) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(script), *(str(arg) for arg in args)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )


def standard_probe(work: Path) -> dict:
    return {
        "generated_at": "2026-10-27T10:00:00+08:00",
        "sources": [{
            "path": str(work / "meta_ads.csv"),
            "format": "csv",
            "encoding": "utf-8",
            "delimiter": ",",
            "header_row": 0,
            "columns": [
                {"name": "Date", "date_ratio": 1.0, "numeric_ratio": 0.0, "role_candidates": ["date"]},
                {"name": "Account ID", "date_ratio": 0.0, "numeric_ratio": 0.0, "role_candidates": ["account_id"]},
                {"name": "Campaign ID", "date_ratio": 0.0, "numeric_ratio": 0.0, "role_candidates": ["campaign_id"]},
                {"name": "Ad ID", "date_ratio": 0.0, "numeric_ratio": 0.0, "role_candidates": ["ad_id"]},
                {"name": "Spend", "date_ratio": 0.0, "numeric_ratio": 1.0, "role_candidates": ["spend"]},
                {"name": "Impressions", "date_ratio": 0.0, "numeric_ratio": 1.0, "role_candidates": ["impressions"]},
                {"name": "Clicks", "date_ratio": 0.0, "numeric_ratio": 1.0, "role_candidates": ["clicks"]},
                {"name": "Installs", "date_ratio": 0.0, "numeric_ratio": 1.0, "role_candidates": ["install"]},
            ],
            "grain_candidates": [{"grain": "ad", "confidence": "high"}],
            "date_coverage": {"min": "2026-10-01", "max": "2026-10-21", "distinct_days": 21},
            "distinct_countries": 2,
        }],
    }


def minimal_standard_profile(*, confirmed: bool, pending_p0: bool = False) -> dict:
    profile = yaml.safe_load((ROOT / "examples" / "minimal" / "profile.yaml").read_text(encoding="utf-8"))
    profile["product"] = {"name": "FreshApp", "timezone": "Asia/Shanghai"}
    profile["report_intent"] = {
            "primary_goal": "efficiency",
            "goal_description": "降低获客成本并保持安装规模",
            "period": {
                "type": "weekly" if confirmed else None,
                "confirmed": confirmed,
                "recommendation": {
                    "suggested_type": "weekly",
                    "evidence": ["源数据覆盖三个完整自然周"],
                },
            },
            "comparison": {"mode": "previous_equal_period"},
            "primary_kpis": ["spend", "impressions", "clicks", "install"],
            "dimensions": ["date", "account_id", "campaign_id", "ad_id"],
            "layout_confirmed": not pending_p0,
        }
    profile["ambiguities"] = ([{
            "id": "p0_metric",
            "question": "请确认安装口径",
            "level": "P0",
            "status": "pending",
        }] if pending_p0 else [])
    profile["output"]["artifacts"].insert(0, {
        "kind": "xlsx_workbook",
        "write_mode": "new_file_per_period",
        "target": {"mode": "create_copy", "sample": "standard-report-template.xlsx"},
        "sheets": [
            {
                "id": sheet_id,
                "match": "daily",
                "target_sheet": sheet_name,
                "role": "history",
                "row_key": {"field": "日期", "format": "date:yyyy-mm-dd"},
                "placement": {"mode": "replace_sheet", "same_key": "replace_in_place"},
                "columns": [{"header": "日期", "source": "$report_date", "format": "date:yyyy-mm-dd"}],
            }
            for sheet_id, sheet_name in [
                ("management_summary", "管理摘要"),
                ("detail_data", "明细数据"),
                ("trend_comparison", "趋势环比"),
            ]
        ],
    })
    return profile


class StandardReportTests(unittest.TestCase):
    def test_infer_combines_same_role_metrics_across_multiple_ad_platforms(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            probe = standard_probe(work)
            google = json.loads(json.dumps(probe["sources"][0]))
            google["path"] = str(work / "google_ads.csv")
            google["columns"] = [
                {**column, "name": "Cost" if column["role_candidates"] == ["spend"] else column["name"]}
                for column in google["columns"]
            ]
            probe["sources"].append(google)
            probe_path = work / "probe.json"
            goal_path = work / "business-goal.yaml"
            profile_path = work / "profile.yaml"
            probe_path.write_text(json.dumps(probe, ensure_ascii=False), encoding="utf-8")
            goal_path.write_text(yaml.safe_dump({
                "primary_goal": "efficiency",
                "goal_description": "统一查看 Meta 和 Google 的投放效率",
            }, allow_unicode=True), encoding="utf-8")

            result = run_cli(
                SCRIPTS / "infer_mappings.py",
                "--probe", probe_path,
                "--business-goal", goal_path,
                "--standard-output",
                "--out", profile_path,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            self.assertEqual(profile["metrics"]["spend"]["role"], "spend")
            self.assertEqual(
                set(profile["metrics"]["spend"]["components"]),
                {"meta_ad.spend", "gg_ad.spend"},
            )
            self.assertNotIn("spend_gg_ad", profile["metrics"])

    def test_infer_standard_output_builds_period_recommendation_and_three_sheets(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            probe_path = work / "probe.json"
            goal_path = work / "business-goal.yaml"
            profile_path = work / "profile.yaml"
            probe_path.write_text(json.dumps(standard_probe(work), ensure_ascii=False), encoding="utf-8")
            goal_path.write_text(yaml.safe_dump({
                "primary_goal": "efficiency",
                "goal_description": "降低获客成本并保持安装规模",
            }, allow_unicode=True), encoding="utf-8")

            result = run_cli(
                SCRIPTS / "infer_mappings.py",
                "--probe", probe_path,
                "--business-goal", goal_path,
                "--standard-output",
                "--product", "FreshApp",
                "--out", profile_path,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            intent = profile["report_intent"]
            self.assertFalse(intent["period"]["confirmed"])
            self.assertEqual(intent["period"]["recommendation"]["suggested_type"], "weekly")
            workbook = next(a for a in profile["output"]["artifacts"] if a["kind"] == "xlsx_workbook")
            self.assertEqual(
                [sheet["target_sheet"] for sheet in workbook["sheets"]],
                ["管理摘要", "明细数据", "趋势环比"],
            )

    def test_standard_report_shows_overall_platform_breakdown_and_full_funnel(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            meta_path = work / "meta_ads.csv"
            google_path = work / "google_ads.csv"
            adjust_path = work / "adjust_events.csv"
            common_fields = ["Date", "Account ID", "Campaign ID", "Ad ID", "Impressions", "Clicks", "Installs"]
            with meta_path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=[*common_fields, "Spend"])
                writer.writeheader()
                writer.writerows([
                    {"Date": "2026-10-14", "Account ID": "meta-a", "Campaign ID": "mc", "Ad ID": "ma", "Spend": 100, "Impressions": 10000, "Clicks": 200, "Installs": 20},
                    {"Date": "2026-10-20", "Account ID": "meta-a", "Campaign ID": "mc", "Ad ID": "ma", "Spend": 150, "Impressions": 15000, "Clicks": 300, "Installs": 30},
                ])
            with google_path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=[*common_fields, "Cost"])
                writer.writeheader()
                writer.writerows([
                    {"Date": "2026-10-14", "Account ID": "google-a", "Campaign ID": "gc", "Ad ID": "ga", "Cost": 50, "Impressions": 5000, "Clicks": 100, "Installs": 10},
                    {"Date": "2026-10-20", "Account ID": "google-a", "Campaign ID": "gc", "Ad ID": "ga", "Cost": 100, "Impressions": 10000, "Clicks": 200, "Installs": 20},
                ])
            with adjust_path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=[
                    "Date", "Account ID", "Ad ID", "Registrations", "KYC Success", "First Deposits",
                ])
                writer.writeheader()
                writer.writerows([
                    {"Date": "2026-10-14", "Account ID": "meta-a", "Ad ID": "ma", "Registrations": 10, "KYC Success": 5, "First Deposits": 2},
                    {"Date": "2026-10-14", "Account ID": "google-a", "Ad ID": "ga", "Registrations": 5, "KYC Success": 3, "First Deposits": 1},
                    {"Date": "2026-10-20", "Account ID": "meta-a", "Ad ID": "ma", "Registrations": 15, "KYC Success": 9, "First Deposits": 4},
                    {"Date": "2026-10-20", "Account ID": "google-a", "Ad ID": "ga", "Registrations": 10, "KYC Success": 6, "First Deposits": 3},
                ])
            probe = standard_probe(work)
            google_probe = json.loads(json.dumps(probe["sources"][0]))
            google_probe["path"] = str(google_path)
            google_probe["columns"] = [
                {**column, "name": "Cost" if column["role_candidates"] == ["spend"] else column["name"]}
                for column in google_probe["columns"]
            ]
            probe["sources"][0]["path"] = str(meta_path)
            probe["sources"].append(google_probe)
            probe["sources"].append({
                "path": str(adjust_path),
                "format": "csv", "encoding": "utf-8", "delimiter": ",", "header_row": 0,
                "columns": [
                    {"name": "Date", "date_ratio": 1.0, "numeric_ratio": 0.0, "role_candidates": ["date"]},
                    {"name": "Account ID", "date_ratio": 0.0, "numeric_ratio": 0.0, "role_candidates": ["account_id"]},
                    {"name": "Ad ID", "date_ratio": 0.0, "numeric_ratio": 0.0, "role_candidates": ["ad_id"]},
                    {"name": "Registrations", "date_ratio": 0.0, "numeric_ratio": 1.0, "role_candidates": ["registration"]},
                    {"name": "KYC Success", "date_ratio": 0.0, "numeric_ratio": 1.0, "role_candidates": ["kyc"]},
                    {"name": "First Deposits", "date_ratio": 0.0, "numeric_ratio": 1.0, "role_candidates": ["first_deposit"]},
                ],
                "grain_candidates": [{"grain": "ad", "confidence": "high"}],
                "date_coverage": {"min": "2026-10-14", "max": "2026-10-20", "distinct_days": 2},
                "distinct_countries": 2,
            })
            probe_path = work / "probe.json"
            goal_path = work / "goal.yaml"
            profile_path = work / "profile.yaml"
            template_path = work / "standard-report-template.xlsx"
            output_path = work / "multi-platform.xlsx"
            probe_path.write_text(json.dumps(probe, ensure_ascii=False), encoding="utf-8")
            goal_path.write_text(yaml.safe_dump({
                "primary_goal": "efficiency",
                "goal_description": "统一查看多个广告平台的投放效率",
                "period": {"type": "weekly", "confirmed": True},
                "layout_confirmed": True,
            }, allow_unicode=True), encoding="utf-8")
            inferred = run_cli(
                SCRIPTS / "infer_mappings.py", "--probe", probe_path,
                "--business-goal", goal_path, "--standard-output", "--out", profile_path,
            )
            self.assertEqual(inferred.returncode, 0, inferred.stdout + inferred.stderr)
            profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            for item in profile.get("ambiguities", []):
                if item.get("level") == "P0":
                    item["status"] = "confirmed"
            profile_path.write_text(
                yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8"
            )
            created = run_cli(
                SCRIPTS / "create_standard_template.py", "--profile", profile_path, "--out", template_path,
            )
            self.assertEqual(created.returncode, 0, created.stdout + created.stderr)
            rendered = run_cli(
                SCRIPTS / "render_workbook.py", "--profile", profile_path,
                "--report-type", "standard", "--date", "2026-10-21",
                "--source-file", f"meta_ad={meta_path}",
                "--source-file", f"gg_ad={google_path}",
                "--source-file", f"adj_ad={adjust_path}",
                "--out", output_path,
            )
            self.assertEqual(rendered.returncode, 0, rendered.stdout + rendered.stderr)
            workbook = load_workbook(output_path, data_only=False)
            summary = workbook["管理摘要"]
            headers = {summary.cell(row=7, column=column).value: column for column in range(1, summary.max_column + 1)}
            self.assertTrue({
                "平台", "花费", "注册", "KYC", "首充", "注册→KYC率", "注册成本", "首充成本",
            }.issubset(headers))
            platform = headers["平台"]
            spend, registration = headers["花费"], headers["注册"]
            kyc, first_deposit = headers["KYC"], headers["首充"]
            self.assertEqual(summary.cell(row=8, column=platform).value, "全部平台")
            self.assertEqual(
                [summary.cell(row=8, column=c).value for c in (spend, registration, kyc, first_deposit)],
                [150, 15, 8, 3],
            )
            self.assertEqual(
                [summary.cell(row=9, column=c).value for c in (spend, registration, kyc, first_deposit)],
                [250, 25, 15, 7],
            )
            platform_rows = {
                summary.cell(row=row, column=platform).value: row
                for row in range(8, summary.max_row + 1)
                if summary.cell(row=row, column=platform).value
            }
            self.assertEqual({"全部平台", "meta", "google"}, set(platform_rows))
            self.assertEqual(summary.cell(row=platform_rows["meta"] + 1, column=spend).value, 150)
            self.assertEqual(summary.cell(row=platform_rows["google"] + 1, column=spend).value, 100)
            detail = workbook["明细数据"]
            detail_headers = {detail.cell(row=3, column=column).value: column for column in range(1, detail.max_column + 1)}
            self.assertTrue({"平台", "注册", "KYC", "首充", "注册→KYC率"}.issubset(detail_headers))
            self.assertEqual(
                {detail.cell(row=row, column=detail_headers["平台"]).value for row in range(4, detail.max_row + 1)},
                {"meta", "google"},
            )
            self.assertEqual(
                {str(detail.cell(row=row, column=detail_headers["日期"]).value)[:10] for row in range(4, detail.max_row + 1)},
                {"2026-10-20"},
            )
            trend = workbook["趋势环比"]
            trend_headers = {trend.cell(row=3, column=column).value: column for column in range(1, trend.max_column + 1)}
            self.assertTrue({"平台", "注册", "KYC", "首充", "注册→KYC率"}.issubset(trend_headers))
            self.assertGreaterEqual(summary._charts[0].anchor._from.row, summary.max_row)
            workbook.close()

    def test_infer_proposes_a_join_for_each_ad_platform(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            probe = standard_probe(work)
            google = json.loads(json.dumps(probe["sources"][0]))
            google["path"] = str(work / "google_ads.csv")
            google["columns"] = [
                {**column, "name": "Cost" if column["role_candidates"] == ["spend"] else column["name"]}
                for column in google["columns"]
            ]
            adjust = {
                "path": str(work / "adjust_events.csv"),
                "format": "csv", "encoding": "utf-8", "delimiter": ",", "header_row": 0,
                "columns": [
                    {"name": "Date", "date_ratio": 1.0, "numeric_ratio": 0.0, "role_candidates": ["date"]},
                    {"name": "Account ID", "date_ratio": 0.0, "numeric_ratio": 0.0, "role_candidates": ["account_id"]},
                    {"name": "Ad ID", "date_ratio": 0.0, "numeric_ratio": 0.0, "role_candidates": ["ad_id"]},
                    {"name": "Registrations", "date_ratio": 0.0, "numeric_ratio": 1.0, "role_candidates": ["registration"]},
                ],
                "grain_candidates": [{"grain": "ad", "confidence": "high"}],
                "date_coverage": {"min": "2026-10-01", "max": "2026-10-21", "distinct_days": 21},
                "distinct_countries": 2,
            }
            probe["sources"].extend([google, adjust])
            probe_path = work / "probe.json"
            goal_path = work / "goal.yaml"
            profile_path = work / "profile.yaml"
            probe_path.write_text(json.dumps(probe, ensure_ascii=False), encoding="utf-8")
            goal_path.write_text(yaml.safe_dump({
                "primary_goal": "acquisition", "goal_description": "提升多平台注册量",
            }, allow_unicode=True), encoding="utf-8")
            result = run_cli(
                SCRIPTS / "infer_mappings.py", "--probe", probe_path,
                "--business-goal", goal_path, "--standard-output", "--out", profile_path,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            self.assertEqual(
                {join["left"] for join in profile["source_plan"]["joins"]},
                {"meta_ad", "gg_ad"},
            )

    def test_template_generation_blocks_when_period_is_unconfirmed(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            profile_path = work / "profile.yaml"
            output_path = work / "report.xlsx"
            profile_path.write_text(yaml.safe_dump(
                minimal_standard_profile(confirmed=False), allow_unicode=True, sort_keys=False
            ), encoding="utf-8")

            result = run_cli(
                SCRIPTS / "create_standard_template.py",
                "--profile", profile_path,
                "--out", output_path,
            )

            self.assertNotEqual(result.returncode, 0)
            self.assertIn("report_intent.period", result.stderr)
            self.assertFalse(output_path.exists())
            self.assertFalse((work / "report_DRAFT.xlsx").exists())

    def test_vague_business_goal_becomes_p0_with_actionable_options(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            probe_path = work / "probe.json"
            goal_path = work / "business-goal.yaml"
            profile_path = work / "profile.yaml"
            probe_path.write_text(json.dumps(standard_probe(work), ensure_ascii=False), encoding="utf-8")
            goal_path.write_text("{}\n", encoding="utf-8")
            result = run_cli(
                SCRIPTS / "infer_mappings.py",
                "--probe", probe_path,
                "--business-goal", goal_path,
                "--standard-output",
                "--out", profile_path,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            issue = next(item for item in profile["ambiguities"] if item["id"] == "p0_business_goal")
            self.assertEqual(issue["status"], "pending")
            self.assertIn("acquisition", issue["suggested"])
            self.assertIn("efficiency", issue["suggested"])
            self.assertIn("revenue", issue["suggested"])

    def test_validator_blocks_unconfirmed_period_and_accepts_confirmed_standard_profile(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            probe_path = work / "probe.json"
            goal_path = work / "business-goal.yaml"
            profile_path = work / "profile.yaml"
            probe_path.write_text(json.dumps(standard_probe(work), ensure_ascii=False), encoding="utf-8")
            goal_path.write_text(yaml.safe_dump({
                "primary_goal": "efficiency",
                "goal_description": "降低获客成本并保持安装规模",
            }, allow_unicode=True), encoding="utf-8")
            inferred = run_cli(
                SCRIPTS / "infer_mappings.py",
                "--probe", probe_path,
                "--business-goal", goal_path,
                "--standard-output",
                "--product", "FreshApp",
                "--out", profile_path,
            )
            self.assertEqual(inferred.returncode, 0, inferred.stdout + inferred.stderr)

            blocked = run_cli(SCRIPTS / "validate_profile.py", profile_path, "--json")
            self.assertNotEqual(blocked.returncode, 0)
            self.assertIn("UNCONFIRMED_PERIOD", blocked.stdout)

            profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            profile["report_intent"]["period"]["type"] = "weekly"
            profile["report_intent"]["period"]["confirmed"] = True
            profile["report_intent"]["layout_confirmed"] = True
            profile["ambiguities"] = [
                item for item in profile.get("ambiguities", []) if item.get("level") != "P0"
            ]
            profile_path.write_text(yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8")

            accepted = run_cli(SCRIPTS / "validate_profile.py", profile_path, "--json")
            self.assertEqual(accepted.returncode, 0, accepted.stdout + accepted.stderr)

    def test_pending_p0_can_only_generate_a_marked_draft(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            profile_path = work / "profile.yaml"
            output_path = work / "report.xlsx"
            profile_path.write_text(yaml.safe_dump(
                minimal_standard_profile(confirmed=True, pending_p0=True), allow_unicode=True, sort_keys=False
            ), encoding="utf-8")

            formal = run_cli(
                SCRIPTS / "create_standard_template.py",
                "--profile", profile_path,
                "--out", output_path,
            )
            self.assertNotEqual(formal.returncode, 0)
            self.assertFalse(output_path.exists())

            draft = run_cli(
                SCRIPTS / "create_standard_template.py",
                "--profile", profile_path,
                "--out", output_path,
                "--draft",
            )
            draft_path = work / "report_DRAFT.xlsx"
            self.assertEqual(draft.returncode, 0, draft.stdout + draft.stderr)
            self.assertTrue(draft_path.exists())
            workbook = load_workbook(draft_path, data_only=False)
            self.assertEqual(workbook.sheetnames, ["管理摘要", "明细数据", "趋势环比"])
            for sheet in workbook.worksheets:
                self.assertIn("DRAFT", str(sheet["A1"].value))
                self.assertEqual(sheet["A1"].fill.fgColor.rgb, "00FFF2CC")
            workbook.close()

    def test_confirmed_profile_generates_exactly_three_formal_sheets(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            profile_path = work / "profile.yaml"
            output_path = work / "report.xlsx"
            profile_path.write_text(yaml.safe_dump(
                minimal_standard_profile(confirmed=True), allow_unicode=True, sort_keys=False
            ), encoding="utf-8")

            result = run_cli(
                SCRIPTS / "create_standard_template.py",
                "--profile", profile_path,
                "--out", output_path,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            workbook = load_workbook(output_path, data_only=False)
            self.assertEqual(workbook.sheetnames, ["管理摘要", "明细数据", "趋势环比"])
            self.assertNotIn("DRAFT", str(workbook["管理摘要"]["A1"].value))
            self.assertEqual(workbook["管理摘要"]["A2"].value, "FreshApp")
            self.assertTrue(workbook["管理摘要"]._charts)
            workbook.close()

    def test_template_generator_rejects_invalid_profile_and_does_not_overwrite(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            valid_path = work / "valid.yaml"
            invalid_path = work / "invalid.yaml"
            output_path = work / "report.xlsx"
            valid = minimal_standard_profile(confirmed=True)
            valid_path.write_text(yaml.safe_dump(valid, allow_unicode=True, sort_keys=False), encoding="utf-8")
            invalid = yaml.safe_load(yaml.safe_dump(valid))
            invalid.pop("product")
            invalid_path.write_text(yaml.safe_dump(invalid, allow_unicode=True, sort_keys=False), encoding="utf-8")

            rejected = run_cli(
                SCRIPTS / "create_standard_template.py",
                "--profile", invalid_path,
                "--out", output_path,
            )
            self.assertNotEqual(rejected.returncode, 0)
            self.assertFalse(output_path.exists())

            created = run_cli(
                SCRIPTS / "create_standard_template.py",
                "--profile", valid_path,
                "--out", output_path,
            )
            self.assertEqual(created.returncode, 0, created.stdout + created.stderr)
            original = output_path.read_bytes()
            refused = run_cli(
                SCRIPTS / "create_standard_template.py",
                "--profile", valid_path,
                "--out", output_path,
            )
            self.assertNotEqual(refused.returncode, 0)
            self.assertEqual(output_path.read_bytes(), original)

    def test_end_to_end_standard_report_has_correct_previous_current_and_change(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            source_path = work / "meta_ads.csv"
            with source_path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=[
                    "Date", "Account ID", "Campaign ID", "Ad ID",
                    "Spend", "Impressions", "Clicks", "Installs",
                ])
                writer.writeheader()
                writer.writerows([
                    {"Date": "2026-10-13", "Account ID": "a1", "Campaign ID": "c1", "Ad ID": "ad1", "Spend": 40, "Impressions": 4000, "Clicks": 80, "Installs": 8},
                    {"Date": "2026-10-14", "Account ID": "a1", "Campaign ID": "c1", "Ad ID": "ad1", "Spend": 60, "Impressions": 6000, "Clicks": 120, "Installs": 12},
                    {"Date": "2026-10-19", "Account ID": "a1", "Campaign ID": "c1", "Ad ID": "ad1", "Spend": 100, "Impressions": 10000, "Clicks": 200, "Installs": 20},
                    {"Date": "2026-10-20", "Account ID": "a1", "Campaign ID": "c1", "Ad ID": "ad1", "Spend": 50, "Impressions": 5000, "Clicks": 100, "Installs": 10},
                ])
            probe = standard_probe(work)
            probe_path = work / "probe.json"
            goal_path = work / "business-goal.yaml"
            profile_path = work / "profile.yaml"
            template_path = work / "standard-report-template.xlsx"
            output_path = work / "standard-report.xlsx"
            probe_path.write_text(json.dumps(probe, ensure_ascii=False), encoding="utf-8")
            goal_path.write_text(yaml.safe_dump({
                "primary_goal": "efficiency",
                "goal_description": "降低获客成本并保持安装规模",
                "period": {"type": "weekly", "confirmed": True},
                "layout_confirmed": True,
            }, allow_unicode=True), encoding="utf-8")

            inferred = run_cli(
                SCRIPTS / "infer_mappings.py",
                "--probe", probe_path,
                "--business-goal", goal_path,
                "--standard-output",
                "--product", "FreshApp",
                "--out", profile_path,
            )
            self.assertEqual(inferred.returncode, 0, inferred.stdout + inferred.stderr)
            created = run_cli(
                SCRIPTS / "create_standard_template.py",
                "--profile", profile_path,
                "--out", template_path,
            )
            self.assertEqual(created.returncode, 0, created.stdout + created.stderr)

            rendered = run_cli(
                SCRIPTS / "render_workbook.py",
                "--profile", profile_path,
                "--report-type", "standard",
                "--date", "2026-10-21",
                "--source-file", f"meta_ad={source_path}",
                "--out", output_path,
            )
            self.assertEqual(rendered.returncode, 0, rendered.stdout + rendered.stderr)
            workbook = load_workbook(output_path, data_only=False)
            self.assertEqual(workbook.sheetnames, ["管理摘要", "明细数据", "趋势环比"])
            summary = workbook["管理摘要"]
            headers = {summary.cell(row=7, column=column).value: column for column in range(1, summary.max_column + 1)}
            spend_column = headers["花费"]
            self.assertEqual(summary.cell(row=8, column=spend_column).value, 100)
            self.assertEqual(summary.cell(row=9, column=spend_column).value, 150)
            self.assertEqual(summary.cell(row=10, column=spend_column).value, 50)
            self.assertEqual(summary.cell(row=11, column=spend_column).value, 0.5)
            self.assertEqual(summary.cell(row=11, column=spend_column).number_format, "0.00%")
            self.assertIn("150.00", summary["B5"].value)
            self.assertIn("+50.00%", summary["B5"].value)
            self.assertNotIn("DRAFT", str(summary["A1"].value))
            workbook.close()

            current_only_path = work / "meta_ads_current_only.csv"
            with current_only_path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=[
                    "Date", "Account ID", "Campaign ID", "Ad ID",
                    "Spend", "Impressions", "Clicks", "Installs",
                ])
                writer.writeheader()
                writer.writerow({"Date": "2026-10-19", "Account ID": "a1", "Campaign ID": "c1", "Ad ID": "ad1", "Spend": 100, "Impressions": 10000, "Clicks": 200, "Installs": 20})
            missing_output = work / "standard-report-missing-previous.xlsx"
            missing = run_cli(
                SCRIPTS / "render_workbook.py",
                "--profile", profile_path,
                "--report-type", "standard",
                "--date", "2026-10-21",
                "--source-file", f"meta_ad={current_only_path}",
                "--out", missing_output,
            )
            self.assertEqual(missing.returncode, 0, missing.stdout + missing.stderr)
            missing_book = load_workbook(missing_output, data_only=False)
            missing_summary = missing_book["管理摘要"]
            self.assertIsNone(missing_summary.cell(row=8, column=spend_column).value)
            self.assertEqual(missing_summary.cell(row=10, column=spend_column).value, "N/A")
            self.assertEqual(missing_summary.cell(row=11, column=spend_column).value, "N/A")
            self.assertIn("N/A", missing_summary["B5"].value)
            missing_book.close()

    def test_custom_period_uses_previous_equal_length_window(self) -> None:
        with tempfile.TemporaryDirectory(dir=ROOT / "runtime") as tmp:
            work = Path(tmp)
            source_path = work / "meta_ads.csv"
            with source_path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=[
                    "Date", "Account ID", "Campaign ID", "Ad ID",
                    "Spend", "Impressions", "Clicks", "Installs",
                ])
                writer.writeheader()
                for day, spend in [(7, 10), (8, 20), (9, 30), (10, 40), (11, 50), (12, 60)]:
                    writer.writerow({
                        "Date": f"2026-10-{day:02d}", "Account ID": "a1", "Campaign ID": "c1",
                        "Ad ID": "ad1", "Spend": spend, "Impressions": spend * 100,
                        "Clicks": spend * 2, "Installs": spend // 5,
                    })
            probe_path = work / "probe.json"
            goal_path = work / "business-goal.yaml"
            profile_path = work / "profile.yaml"
            out_dir = work / "out"
            probe_path.write_text(json.dumps(standard_probe(work), ensure_ascii=False), encoding="utf-8")
            goal_path.write_text(yaml.safe_dump({
                "primary_goal": "custom",
                "goal_description": "复盘活动窗口",
                "period": {"type": "custom", "start": "2026-10-10", "end": "2026-10-12", "confirmed": True},
                "layout_confirmed": True,
            }, allow_unicode=True), encoding="utf-8")
            inferred = run_cli(
                SCRIPTS / "infer_mappings.py",
                "--probe", probe_path,
                "--business-goal", goal_path,
                "--standard-output",
                "--product", "FreshApp",
                "--out", profile_path,
            )
            self.assertEqual(inferred.returncode, 0, inferred.stdout + inferred.stderr)
            generated = run_cli(
                SCRIPTS / "generate_report.py",
                "--profile", profile_path,
                "--report-type", "standard",
                "--date", "2026-10-12",
                "--source-file", f"meta_ad={source_path}",
                "--out-dir", out_dir,
            )
            self.assertEqual(generated.returncode, 0, generated.stdout + generated.stderr)
            with (out_dir / "过程表_20261012.csv").open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            totals: dict[str, float] = {}
            for row in rows:
                totals[row["周期"]] = totals.get(row["周期"], 0) + float(row["花费"])
            self.assertEqual(totals["previous_custom"], 60)
            self.assertEqual(totals["current_custom"], 150)


if __name__ == "__main__":
    unittest.main()
