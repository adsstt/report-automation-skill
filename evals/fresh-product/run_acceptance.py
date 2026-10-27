from __future__ import annotations

import csv
import os
import shutil
import subprocess
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
PACKAGE = ROOT.parents[1]
SCRIPTS = PACKAGE / "scripts"
RUNTIME = ROOT / "runtime"


def run(name: str, args: list[object], expected: int = 0) -> str:
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    result = subprocess.run(
        [sys.executable, *(str(arg) for arg in args)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=env,
    )
    output = (result.stdout + result.stderr).strip()
    print(f"[{name}] exit={result.returncode}")
    if output:
        console_encoding = sys.stdout.encoding or "utf-8"
        print(output.encode(console_encoding, errors="replace").decode(console_encoding))
    if result.returncode != expected:
        raise RuntimeError(f"{name}: expected exit {expected}, got {result.returncode}")
    return output


def main() -> int:
    if RUNTIME.exists():
        if RUNTIME.resolve().parent != ROOT.resolve():
            raise RuntimeError("refusing to clean runtime outside scenario root")
        shutil.rmtree(RUNTIME)
    RUNTIME.mkdir(parents=True)

    run("fixtures", [ROOT / "make_fixtures.py"])
    probe = RUNTIME / "probe.json"
    draft = RUNTIME / "profile.draft.yaml"
    out_dir = RUNTIME / "output"
    workbook = RUNTIME / "NovaFit_daily_20261025.xlsx"

    run("probe", [
        SCRIPTS / "probe_sources.py",
        "--source-dir", ROOT / "sources" / "2026-10-24",
        "--source-dir", ROOT / "sources" / "2026-10-25",
        "--glob", "*",
        "--limit", "10",
        "--target", ROOT / "samples" / "target_template.xlsx",
        "--out", probe,
    ])
    run("infer", [SCRIPTS / "infer_mappings.py", "--probe", probe, "--out", draft, "--product", "NovaFit"])
    run("draft gate", [SCRIPTS / "validate_profile.py", draft], expected=1)
    run("confirmed profile", [SCRIPTS / "validate_profile.py", ROOT / "profile.confirmed.yaml"])
    run("format", [
        SCRIPTS / "verify_format.py", "--profile", ROOT / "profile.confirmed.yaml",
        "--target", ROOT / "samples" / "target_template.xlsx",
    ])
    run("generate", [
        SCRIPTS / "generate_report.py", "--profile", ROOT / "profile.confirmed.yaml",
        "--report-type", "daily", "--date", "2026-10-25", "--out-dir", out_dir,
    ])
    run("numeric verification", [
        SCRIPTS / "verify_output.py", "--profile", ROOT / "profile.confirmed.yaml",
        "--out-dir", out_dir, "--date", "2026-10-25",
    ])
    run("shadow diff", [
        SCRIPTS / "shadow_diff.py",
        "--shadow", out_dir / "结果表_20261025.csv",
        "--expected", ROOT / "expected" / "结果表_20261025.csv",
        "--keys", "账户ID,国家代码,事件",
        "--out", RUNTIME / "shadow_diff.md",
    ])
    run("render", [
        SCRIPTS / "render_workbook.py", "--profile", ROOT / "profile.confirmed.yaml",
        "--report-type", "daily", "--date", "2026-10-25",
        "--source-file", f"meta_ads={ROOT / 'sources' / '2026-10-25' / 'meta_ads_20261025.xlsx'}",
        "--source-file", f"adjust_events={ROOT / 'sources' / '2026-10-25' / 'adjust_attribution_20261025.csv'}",
        "--compare-source-file", f"meta_ads={ROOT / 'sources' / '2026-10-24' / 'meta_ads_20261024.xlsx'}",
        "--compare-source-file", f"adjust_events={ROOT / 'sources' / '2026-10-24' / 'adjust_attribution_20261024.csv'}",
        "--out", workbook,
    ])
    run("workbook verification", [
        SCRIPTS / "verify_workbook.py",
        "--template", ROOT / "samples" / "target_template.xlsx",
        "--output", workbook,
        "--report-date", "2026-10-25",
        "--sheet", "Daily Comparison",
        "--expect-sheet", "Daily History",
        "--expect-new-rows", "2",
    ])

    with (out_dir / "过程表_20261025.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    totals = {
        "spend": sum(float(row["花费"]) for row in rows),
        "registration": sum(float(row["注册"]) for row in rows),
        "kyc": sum(float(row["KYC"]) for row in rows),
        "ftd": sum(float(row["FTD_Actual"]) for row in rows),
        "revenue": sum(float(row["收入"]) for row in rows),
    }
    expected = {"spend": 300, "registration": 30, "kyc": 18, "ftd": 8, "revenue": 800}
    if totals != expected:
        raise RuntimeError(f"unexpected core totals: {totals}")

    # 第二条独立验收：新用户没有目标模板，只提供源数据与业务目标。
    standard_probe = RUNTIME / "standard-probe.json"
    standard_draft = RUNTIME / "standard-profile.draft.yaml"
    standard_confirmed = RUNTIME / "standard-profile.confirmed.yaml"
    standard_template = RUNTIME / "standard-report-template.xlsx"
    standard_workbook = RUNTIME / "NovaFit_standard_weekly_20261025.xlsx"
    standard_source = ROOT / "standard-sources" / "novafit_meta_ads.csv"
    run("standard probe", [
        SCRIPTS / "probe_sources.py", "--source-file", standard_source, "--out", standard_probe,
    ])
    run("standard infer", [
        SCRIPTS / "infer_mappings.py", "--probe", standard_probe,
        "--business-goal", ROOT / "business-goal.standard.yaml", "--standard-output",
        "--product", "NovaFit", "--out", standard_draft,
    ])
    run("unconfirmed period gate", [SCRIPTS / "validate_profile.py", standard_draft], expected=1)
    run("unconfirmed period template gate", [
        SCRIPTS / "create_standard_template.py", "--profile", standard_draft, "--out", standard_template,
    ], expected=1)
    if standard_template.exists():
        raise RuntimeError("unconfirmed period unexpectedly created a workbook")

    standard_profile = yaml.safe_load(standard_draft.read_text(encoding="utf-8"))
    standard_profile["report_intent"]["period"]["type"] = "weekly"
    standard_profile["report_intent"]["period"]["confirmed"] = True
    standard_profile["report_intent"]["layout_confirmed"] = True
    for ambiguity in standard_profile.get("ambiguities", []):
        if ambiguity.get("level") == "P0":
            ambiguity["status"] = "confirmed"
    standard_confirmed.write_text(
        yaml.safe_dump(standard_profile, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )
    run("standard confirmed profile", [SCRIPTS / "validate_profile.py", standard_confirmed])
    run("standard template", [
        SCRIPTS / "create_standard_template.py", "--profile", standard_confirmed,
        "--out", standard_template,
    ])
    run("standard render", [
        SCRIPTS / "render_workbook.py", "--profile", standard_confirmed,
        "--report-type", "standard", "--date", "2026-10-25",
        "--source-file", f"meta_ad={standard_source}", "--out", standard_workbook,
    ])
    standard_book = load_workbook(standard_workbook, data_only=False)
    if standard_book.sheetnames != ["管理摘要", "明细数据", "趋势环比"]:
        raise RuntimeError(f"unexpected standard sheets: {standard_book.sheetnames}")
    summary = standard_book["管理摘要"]
    headers = {summary.cell(row=7, column=column).value: column for column in range(1, summary.max_column + 1)}
    spend_column = headers["花费"]
    standard_values = [summary.cell(row=row, column=spend_column).value for row in range(8, 12)]
    standard_book.close()
    if standard_values != [280, 350, 70, 0.25]:
        raise RuntimeError(f"unexpected standard comparison values: {standard_values}")

    (RUNTIME / "acceptance_result.md").write_text(
        "# NovaFit fresh-product acceptance\n\n"
        "- Result: PASS\n"
        "- Template flow: Probe → Infer (blocked draft) → confirmed profile → generate → verify → render → workbook verify\n"
        "- No-template flow: Probe → period recommendation/block → user confirms weekly/P0/layout → standard template → formal three-sheet report\n"
        f"- Core totals: {totals}\n"
        f"- Workbook: `{workbook}`\n"
        f"- Standard comparison (previous/current/absolute/rate): {standard_values}\n"
        f"- Standard workbook: `{standard_workbook}`\n",
        encoding="utf-8",
    )
    print(f"ACCEPTANCE PASS: template={totals}; no-template={standard_values}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
