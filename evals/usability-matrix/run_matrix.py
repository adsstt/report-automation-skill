from __future__ import annotations

import csv
import json
import shutil
import subprocess
import sys
from datetime import date, timedelta
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parent
PACKAGE = ROOT.parents[1]
RUNTIME = ROOT / "runtime"
NPM = shutil.which("npm.cmd") or shutil.which("npm") or "npm"
NODE = shutil.which("node.exe") or shutil.which("node") or "node"


def run(name: str, command: list[object], *, cwd: Path, expected: int = 0) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [str(item) for item in command], cwd=cwd, capture_output=True, text=True,
        encoding="utf-8", errors="replace",
    )
    print(f"[{name}] exit={result.returncode}")
    if result.returncode != expected:
        raise RuntimeError(
            f"{name}: expected {expected}, got {result.returncode}\n{result.stdout}\n{result.stderr}"
        )
    return result


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def make_template(path: Path, sheet_name: str = "Daily") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append(["Date", "Account ID", "Spend", "Registrations"])
    book.save(path)
    book.close()


def source(
    source_id: str, platform: str, role: str, path: Path, schema: list[dict[str, str]],
    *, snapshot: bool = False,
) -> dict:
    return {
        "id": source_id, "platform": platform, "role": role, "grain": "ad",
        "discover": {"dir": str(path.parent), "glob": path.name, "pick": {"strategy": "explicit"}},
        "layout": {"header_row": 0, "sheet": 0, "encoding": "utf-8", "csv_delimiter": ","},
        "date_column": {"raw": "Date", "parse": "%Y-%m-%d"},
        "observed_schema": schema,
        "required_columns": ["date", "account_id", "ad_id"],
        "snapshot": ({"is_snapshot": True, "snapshot_field": "snapshot_date"} if snapshot else {"is_snapshot": False}),
    }


def target_profile(
    product: str, template: Path, ad_source: dict, attribution_source: dict,
    metrics: dict, *, join_key: str = "ad_id",
) -> dict:
    return {
        "profile_version": 1,
        "product": {"name": product, "timezone": "Asia/Shanghai"},
        "accounts": {"include": [], "exclude": []},
        "report_types": [{
            "id": "daily",
            "time_scope": {"mode": "single_period", "period": "day", "anchor": "report_date"},
            "aggregation": {"grain": "ad", "group_by": ["account_id"], "total_row": True},
            "comparison": {"mode": "none"},
        }],
        "sources": [ad_source, attribution_source],
        "source_plan": {
            "analysis_grain": "ad", "grains_provided": ["ad"],
            "joins": [{
                "id": f"{ad_source['id']}_x_{attribution_source['id']}",
                "left": ad_source["id"], "right": attribution_source["id"], "type": "left",
                "keys": ["date", "account_id", join_key], "enforce_grain": "ad",
            }],
            "additive_policy": {"forbid_cross_grain_sum": True, "additive_measures": list(metrics)},
            "dedupe": [{"source": attribution_source["id"], "key": ["date", "ad_id"], "keep": "latest_snapshot"}],
            "unmatched_policy": {"keep_left": True, "report": True, "label": "Unmatched"},
        },
        "metrics": metrics,
        "derived": {},
        "output": {
            "artifacts": [
                {
                    "kind": "xlsx_workbook", "write_mode": "append_to_sample",
                    "target": {"mode": "append_to_sample", "sample": str(template)},
                    "sheets": [{
                        "id": "daily", "match": "daily", "target_sheet": "Daily", "role": "history",
                        "row_key": {"field": "Date", "format": "date:yyyy-mm-dd"},
                        "placement": {"mode": "replace_sheet", "same_key": "replace_in_place"},
                        "column_match": "by_header", "group_by": ["account_id"],
                        "columns": [
                            {"header": "Date", "source": "$report_date", "format": "date:yyyy-mm-dd"},
                            {"header": "Account ID", "source": "group.account_id", "format": "text"},
                            {"header": "Spend", "source": "metrics.spend", "format": "number:2"},
                            {"header": "Registrations", "source": "metrics.registration", "format": "number:0"},
                        ],
                    }],
                },
                {"kind": "csv", "role": "process_table", "path": "output/process_{date}.csv", "grain": "ad"},
                {"kind": "csv", "role": "result_table", "path": "output/result_{date}.csv", "group_by": ["account_id"]},
            ]
        },
        "validation": {
            "default_on_failure": "block",
            "checks": [
                {"id": "total_reconcile", "rule": "totals match", "tolerance": 0.01, "level": "blocking"},
                {"id": "row_count", "rule": "row count match", "level": "blocking"},
            ],
        },
        "ambiguities": [],
        "origin": {"type": "from_scratch", "verified": ["probe", "shadow_diff"], "unverified_items": 0},
    }


def install_fresh() -> Path:
    pack_dir = RUNTIME / "package"
    consumer = RUNTIME / "consumer"
    fake_home = RUNTIME / "fresh-home"
    pack_dir.mkdir(parents=True)
    run("npm pack", [NPM, "pack", "--pack-destination", pack_dir], cwd=PACKAGE)
    tarball = next(pack_dir.glob("report-automation-skill-*.tgz"))
    consumer.mkdir(parents=True)
    run(
        "npm install tarball",
        [NPM, "install", "--prefix", consumer, tarball, "--ignore-scripts", "--no-audit", "--no-fund"],
        cwd=PACKAGE,
    )
    package_root = consumer / "node_modules" / "report-automation-skill"
    code = (
        "import { pathToFileURL } from 'node:url';"
        f"const entry=await import(pathToFileURL({json.dumps(str(package_root / 'index.js'))}).href);"
        f"entry.syncSkill({{homeDir:{json.dumps(str(fake_home))},sourceDir:{json.dumps(str(package_root))}}});"
    )
    run("DSH sync", [NODE, "--input-type=module", "-e", code], cwd=PACKAGE)
    installed = fake_home / ".agents" / "skills" / "report-automation"
    required = [
        installed / "SKILL.md", installed / "scripts" / "infer_mappings.py",
        installed / "scripts" / "create_standard_template.py",
    ]
    if not all(path.exists() for path in required):
        raise RuntimeError(f"fresh install missing files: {[str(path) for path in required if not path.exists()]}")
    return installed


def scenario_atlaspay(installed: Path) -> dict:
    work = RUNTIME / "atlaspay"
    work.mkdir()
    ads = work / "tiktok_ads.csv"
    af = work / "appsflyer_events.csv"
    template = work / "target.xlsx"
    write_csv(ads, [{
        "Date": "2026-10-27", "Account ID": "tt-a1", "Ad ID": "tt-ad1",
        "Spend": 100, "Impressions": 10000, "Clicks": 500, "Installs": 50,
    }])
    write_csv(af, [{"Date": "2026-10-27", "Account ID": "tt-a1", "Ad ID": "tt-ad1", "Registrations": 20}])
    make_template(template)
    probe = work / "probe.json"
    draft = work / "draft.yaml"
    run("AtlasPay probe", [
        sys.executable, installed / "scripts" / "probe_sources.py",
        "--source-file", ads, "--source-file", af, "--target", template, "--out", probe,
    ], cwd=work)
    run("AtlasPay infer", [
        sys.executable, installed / "scripts" / "infer_mappings.py", "--probe", probe,
        "--product", "AtlasPay", "--out", draft,
    ], cwd=work)
    draft_profile = yaml.safe_load(draft.read_text(encoding="utf-8"))
    if not any(item.get("id") == "p0_join_keys" for item in draft_profile.get("ambiguities", [])):
        raise RuntimeError("AtlasPay did not ask the knowledgeable user to confirm join keys")

    ad_source = source("tt_ad", "tiktok", "ad_platform", ads, [
        {"raw": "Date", "canonical": "date", "dtype": "date"},
        {"raw": "Account ID", "canonical": "account_id", "dtype": "text"},
        {"raw": "Ad ID", "canonical": "ad_id", "dtype": "text"},
        {"raw": "Spend", "canonical": "spend", "dtype": "number"},
        {"raw": "Impressions", "canonical": "impressions", "dtype": "number"},
        {"raw": "Clicks", "canonical": "clicks", "dtype": "number"},
        {"raw": "Installs", "canonical": "install", "dtype": "number"},
    ])
    af_source = source("af_ad", "appsflyer", "attribution", af, [
        {"raw": "Date", "canonical": "date", "dtype": "date"},
        {"raw": "Account ID", "canonical": "account_id", "dtype": "text"},
        {"raw": "Ad ID", "canonical": "ad_id", "dtype": "text"},
        {"raw": "Registrations", "canonical": "registration", "dtype": "number"},
    ], snapshot=True)
    profile = target_profile("AtlasPay", template, ad_source, af_source, {
        "spend": {"role": "spend", "from": "tt_ad.spend", "agg": "sum"},
        "registration": {"role": "registration", "from": "af_ad.registration", "agg": "sum"},
    })
    profile_path = work / "profile.confirmed.yaml"
    profile_path.write_text(yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
    run("AtlasPay validate", [sys.executable, installed / "scripts" / "validate_profile.py", profile_path], cwd=work)
    output = work / "AtlasPay_daily.xlsx"
    run("AtlasPay render", [
        sys.executable, installed / "scripts" / "render_workbook.py", "--profile", profile_path,
        "--report-type", "daily", "--date", "2026-10-27",
        "--source-file", f"tt_ad={ads}", "--source-file", f"af_ad={af}", "--out", output,
    ], cwd=work)
    book = load_workbook(output, data_only=True)
    values = [book["Daily"]["C2"].value, book["Daily"]["D2"].value]
    book.close()
    if values != [100, 20]:
        raise RuntimeError(f"AtlasPay totals mismatch: {values}")
    return {"result": "PASS", "questions": len(draft_profile.get("ambiguities", [])), "totals": values}


def scenario_bloomcart(installed: Path) -> dict:
    work = RUNTIME / "bloomcart"
    work.mkdir()
    meta, google, adjust = work / "meta_ads.csv", work / "google_ads.csv", work / "adjust_events.csv"
    meta_rows, google_rows, adjust_rows = [], [], []
    start = date(2026, 10, 12)
    for offset in range(14):
        day = (start + timedelta(days=offset)).isoformat()
        current = offset >= 7
        meta_rows.append({"Date": day, "Account ID": "meta-a", "Ad ID": "meta-ad", "Spend": 20 if current else 10, "Impressions": 2000, "Clicks": 100, "Installs": 10})
        google_rows.append({"Date": day, "Account ID": "google-a", "Ad ID": "google-ad", "Cost": 15 if current else 5, "Impressions": 1500, "Clicks": 75, "Installs": 8})
        adjust_rows.extend([
            {"Date": day, "Account ID": "meta-a", "Ad ID": "meta-ad", "Registrations": 2 if current else 1},
            {"Date": day, "Account ID": "google-a", "Ad ID": "google-ad", "Registrations": 3 if current else 2},
        ])
    write_csv(meta, meta_rows)
    write_csv(google, google_rows)
    write_csv(adjust, adjust_rows)
    probe = work / "probe.json"
    vague_goal = work / "goal-vague.yaml"
    vague_profile = work / "profile-vague.yaml"
    vague_goal.write_text("{}\n", encoding="utf-8")
    run("BloomCart probe", [
        sys.executable, installed / "scripts" / "probe_sources.py", "--source-file", meta,
        "--source-file", google, "--source-file", adjust, "--out", probe,
    ], cwd=work)
    run("BloomCart vague infer", [
        sys.executable, installed / "scripts" / "infer_mappings.py", "--probe", probe,
        "--business-goal", vague_goal, "--standard-output", "--product", "BloomCart", "--out", vague_profile,
    ], cwd=work)
    vague = yaml.safe_load(vague_profile.read_text(encoding="utf-8"))
    p0_ids = {item["id"] for item in vague.get("ambiguities", []) if item.get("level") == "P0"}
    if "p0_business_goal" not in p0_ids or vague["report_intent"]["period"]["confirmed"]:
        raise RuntimeError("BloomCart novice onboarding did not block vague goal/unconfirmed period")
    run("BloomCart template gate", [
        sys.executable, installed / "scripts" / "create_standard_template.py",
        "--profile", vague_profile, "--out", work / "must-not-exist.xlsx", "--draft",
    ], cwd=work, expected=1)

    goal = work / "goal-confirmed.yaml"
    profile_path = work / "profile.confirmed.yaml"
    goal.write_text(yaml.safe_dump({
        "primary_goal": "efficiency",
        "goal_description": "统一查看两个广告平台的投放效率",
        "period": {"type": "weekly", "confirmed": True},
        "layout_confirmed": True,
    }, allow_unicode=True), encoding="utf-8")
    run("BloomCart confirmed infer", [
        sys.executable, installed / "scripts" / "infer_mappings.py", "--probe", probe,
        "--business-goal", goal, "--standard-output", "--product", "BloomCart", "--out", profile_path,
    ], cwd=work)
    profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
    for item in profile.get("ambiguities", []):
        if item.get("level") == "P0":
            item["status"] = "confirmed"
    profile_path.write_text(yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
    run("BloomCart validate", [sys.executable, installed / "scripts" / "validate_profile.py", profile_path], cwd=work)
    template = work / "standard-report-template.xlsx"
    report = work / "BloomCart_weekly.xlsx"
    run("BloomCart template", [
        sys.executable, installed / "scripts" / "create_standard_template.py",
        "--profile", profile_path, "--out", template,
    ], cwd=work)
    run("BloomCart render", [
        sys.executable, installed / "scripts" / "render_workbook.py", "--profile", profile_path,
        "--report-type", "standard", "--date", "2026-10-25",
        "--source-file", f"meta_ad={meta}", "--source-file", f"gg_ad={google}",
        "--source-file", f"adj_ad={adjust}", "--out", report,
    ], cwd=work)
    book = load_workbook(report, data_only=False)
    if book.sheetnames != ["管理摘要", "明细数据", "趋势环比"]:
        raise RuntimeError(f"BloomCart sheet mismatch: {book.sheetnames}")
    summary = book["管理摘要"]
    headers = {summary.cell(7, column).value: column for column in range(1, summary.max_column + 1)}
    values = {
        "previous_spend": summary.cell(8, headers["花费"]).value,
        "current_spend": summary.cell(9, headers["花费"]).value,
        "previous_registration": summary.cell(8, headers["注册"]).value,
        "current_registration": summary.cell(9, headers["注册"]).value,
    }
    book.close()
    expected = {"previous_spend": 105, "current_spend": 245, "previous_registration": 21, "current_registration": 35}
    if values != expected:
        raise RuntimeError(f"BloomCart totals mismatch: {values}")
    return {"result": "PASS", "initial_p0": len(p0_ids), "totals": values}


def scenario_lumagame(installed: Path) -> dict:
    work = RUNTIME / "lumagame"
    work.mkdir()
    ads, af, template = work / "google_ads.csv", work / "appsflyer_events.csv", work / "target.xlsx"
    write_csv(ads, [{"Date": "2026-10-27", "Account ID": "g-a1", "Ad ID": "g-ad1", "Cost": 80, "Impressions": 8000, "Clicks": 400, "Installs": 40}])
    write_csv(af, [{"Date": "2026-10-27", "Account ID": "g-a1", "Ad ID": "g-ad1", "Registrations Activity": 7, "Registrations LTV": 70}])
    make_template(template)
    probe, draft = work / "probe.json", work / "draft.yaml"
    run("LumaGame probe", [
        sys.executable, installed / "scripts" / "probe_sources.py", "--source-file", ads,
        "--source-file", af, "--target", template, "--out", probe,
    ], cwd=work)
    run("LumaGame infer", [
        sys.executable, installed / "scripts" / "infer_mappings.py", "--probe", probe,
        "--product", "LumaGame", "--out", draft,
    ], cwd=work)
    draft_profile = yaml.safe_load(draft.read_text(encoding="utf-8"))
    candidate_issue = next(
        (item for item in draft_profile.get("ambiguities", []) if "registration" in item.get("id", "") and "candidate" in item.get("id", "")),
        None,
    )
    if candidate_issue is None:
        raise RuntimeError("LumaGame did not block ambiguous Activity/LTV registration fields")
    run("LumaGame draft gate", [sys.executable, installed / "scripts" / "validate_profile.py", draft], cwd=work, expected=1)

    ad_source = source("gg_ad", "google", "ad_platform", ads, [
        {"raw": "Date", "canonical": "date", "dtype": "date"},
        {"raw": "Account ID", "canonical": "account_id", "dtype": "text"},
        {"raw": "Ad ID", "canonical": "ad_id", "dtype": "text"},
        {"raw": "Cost", "canonical": "spend", "dtype": "number"},
    ])
    af_source = source("af_ad", "appsflyer", "attribution", af, [
        {"raw": "Date", "canonical": "date", "dtype": "date"},
        {"raw": "Account ID", "canonical": "account_id", "dtype": "text"},
        {"raw": "Ad ID", "canonical": "ad_id", "dtype": "text"},
    ], snapshot=True)
    registration = {
        "role": "registration",
        "candidates": [
            {"id": "activity", "source": "af_ad", "field": "Registrations Activity", "semantics": {"entity": "unique_user", "window": "daily", "basis": "activity"}, "status": "confirmed"},
            {"id": "ltv", "source": "af_ad", "field": "Registrations LTV", "semantics": {"entity": "unique_user", "window": "ltv", "basis": "ltv"}, "status": "rejected"},
        ],
        "chosen": "activity", "confirm_required": True,
    }
    profile = target_profile("LumaGame", template, ad_source, af_source, {
        "spend": {"role": "spend", "from": "gg_ad.spend", "agg": "sum"},
        "registration": registration,
    })
    profile_path = work / "profile.confirmed.yaml"
    profile_path.write_text(yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
    run("LumaGame validate", [sys.executable, installed / "scripts" / "validate_profile.py", profile_path], cwd=work)
    output = work / "LumaGame_daily.xlsx"
    run("LumaGame render", [
        sys.executable, installed / "scripts" / "render_workbook.py", "--profile", profile_path,
        "--report-type", "daily", "--date", "2026-10-27",
        "--source-file", f"gg_ad={ads}", "--source-file", f"af_ad={af}", "--out", output,
    ], cwd=work)
    book = load_workbook(output, data_only=True)
    values = [book["Daily"]["C2"].value, book["Daily"]["D2"].value]
    book.close()
    if values != [80, 7]:
        raise RuntimeError(f"LumaGame used the wrong registration definition: {values}")
    return {"result": "PASS", "blocked_candidate": candidate_issue["id"], "totals": values}


def main() -> int:
    if RUNTIME.exists():
        if RUNTIME.resolve().parent != ROOT.resolve():
            raise RuntimeError("refusing to clean runtime outside usability matrix")
        shutil.rmtree(RUNTIME)
    RUNTIME.mkdir(parents=True)
    installed = install_fresh()
    results = {
        "fresh_install": {"result": "PASS", "path": str(installed)},
        "AtlasPay": scenario_atlaspay(installed),
        "BloomCart": scenario_bloomcart(installed),
        "LumaGame": scenario_lumagame(installed),
    }
    lines = ["# Report automation usability matrix", "", f"- Fresh install: `{installed}`", ""]
    for name, result in results.items():
        lines.append(f"## {name}")
        lines.append("")
        lines.append("```json")
        lines.append(json.dumps(result, ensure_ascii=False, indent=2))
        lines.append("```")
        lines.append("")
    result_path = RUNTIME / "usability_result.md"
    result_path.write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps(results, ensure_ascii=False, indent=2))
    print(f"USABILITY MATRIX PASS: {result_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
