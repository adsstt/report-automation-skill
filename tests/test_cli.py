from __future__ import annotations

import csv
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
MINIMAL_PROFILE = ROOT / "examples" / "minimal" / "profile.yaml"


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def run_cli(*args: object) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    return subprocess.run(
        [sys.executable, *(str(arg) for arg in args)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=env,
    )


class GenerateReportCliTests(unittest.TestCase):
    def test_multi_period_report_keeps_previous_and_current_totals_separate(self) -> None:
        test_tmp = ROOT / "runtime" / "test-tmp"
        test_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_tmp) as tmp:
            work = Path(tmp)
            profile = yaml.safe_load(MINIMAL_PROFILE.read_text(encoding="utf-8"))
            source = profile["sources"][0]
            source["id"] = "ads"
            source["platform"] = "meta"
            source["date_column"] = {"raw": "Day", "parse": "%Y-%m-%d"}
            source["observed_schema"] = [
                {"raw": "Day", "canonical": "date", "dtype": "date"},
                {"raw": "Account", "canonical": "account_id", "dtype": "text"},
                {"raw": "Ad", "canonical": "ad_id", "dtype": "text"},
                {"raw": "Spend", "canonical": "spend", "dtype": "number"},
            ]
            source["required_columns"] = ["date", "account_id", "ad_id", "spend"]
            profile["sources"] = [source]
            profile["source_plan"] = {
                "analysis_grain": "ad", "grains_provided": ["ad"], "joins": [],
                "additive_policy": {"forbid_cross_grain_sum": True, "additive_measures": ["spend"]},
            }
            profile["metrics"] = {"spend": {"role": "spend", "from": "ads.spend", "agg": "sum"}}
            profile["derived"] = {}
            profile["report_types"] = [{
                "id": "two_day",
                "time_scope": {
                    "mode": "multi_period", "period": "day", "anchor": "report_date",
                    "windows": ["previous_day", "current_day"],
                },
                "aggregation": {"grain": "ad", "group_by": ["account_id"], "total_row": False},
                "comparison": {"mode": "previous_period", "base": "previous_day", "target": "current_day"},
            }]
            profile["output"]["artifacts"][0]["grain"] = "ad"
            profile["output"]["artifacts"][1]["group_by"] = ["账户ID"]
            profile_path = work / "profile.yaml"
            profile_path.write_text(yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
            source_path = work / "ads.csv"
            write_csv(source_path, [
                {"Day": "2026-10-24", "Account": "acct-1", "Ad": "ad-1", "Spend": 10},
                {"Day": "2026-10-25", "Account": "acct-1", "Ad": "ad-1", "Spend": 20},
            ])
            out_dir = work / "out"
            result = run_cli(
                SCRIPTS / "generate_report.py", "--profile", profile_path,
                "--report-type", "two_day", "--date", "2026-10-25",
                "--source-file", f"ads={source_path}", "--out-dir", out_dir,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            with (out_dir / "结果表_20261025.csv").open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(
                [(row["周期"], row["花费"]) for row in rows],
                [("previous_day", "10"), ("current_day", "20")],
            )

    def test_promoted_example_is_created_as_blocked_draft(self) -> None:
        test_tmp = ROOT / "runtime" / "test-tmp"
        test_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_tmp) as tmp:
            promoted = Path(tmp) / "profile.yaml"
            result = run_cli(
                SCRIPTS / "promote_example.py",
                "--example", MINIMAL_PROFILE,
                "--out", promoted,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            profile = yaml.safe_load(promoted.read_text(encoding="utf-8"))
            self.assertEqual(profile["origin"]["type"], "from_example")
            validation = run_cli(SCRIPTS / "validate_profile.py", promoted)
            self.assertEqual(validation.returncode, 1)
            self.assertIn("MISSING_VERIFICATION", validation.stdout)

    def test_example_profile_cannot_generate_official_report(self) -> None:
        result = run_cli(
            SCRIPTS / "generate_report.py",
            "--profile", MINIMAL_PROFILE,
            "--report-type", "daily",
            "--date", "2026-10-25",
            "--out-dir", ROOT / "runtime" / "should-not-exist",
        )
        self.assertEqual(result.returncode, 1)
        self.assertIn("examples", result.stderr)

    def test_infer_does_not_default_unknown_platform_to_tiktok(self) -> None:
        test_tmp = ROOT / "runtime" / "test-tmp"
        test_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_tmp) as tmp:
            work = Path(tmp)
            probe_path = work / "probe.json"
            profile_path = work / "profile.yaml"
            probe_path.write_text(json.dumps({
                "sources": [{
                    "path": str(work / "mystery.csv"),
                    "format": "csv",
                    "encoding": "utf-8",
                    "header_row": 0,
                    "columns": [
                        {"name": "Date", "date_ratio": 1.0, "numeric_ratio": 0.0, "role_candidates": ["date"]},
                        {"name": "Ad ID", "date_ratio": 0.0, "numeric_ratio": 0.0, "role_candidates": ["ad_id"]},
                        {"name": "Spend", "date_ratio": 0.0, "numeric_ratio": 1.0, "role_candidates": ["spend"]},
                    ],
                    "grain_candidates": [{"grain": "ad", "confidence": "high"}],
                    "date_coverage": {"distinct_days": 1},
                }],
            }), encoding="utf-8")

            result = run_cli(
                SCRIPTS / "infer_mappings.py",
                "--probe", probe_path,
                "--out", profile_path,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            self.assertEqual(profile["sources"][0]["platform"], "unknown")
            pending = {item["id"] for item in profile["ambiguities"] if item["status"] == "pending"}
            self.assertIn("p0_source_platform", pending)
            metric_evidence = next(
                item for item in profile["inference"]["candidates"]
                if item["subject"].startswith("metrics.spend")
            )
            self.assertEqual(metric_evidence["confidence"], "low")

    def test_workbook_comparison_writes_metric_delta_ratio(self) -> None:
        test_tmp = ROOT / "runtime" / "test-tmp"
        test_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_tmp) as tmp:
            work = Path(tmp)
            profile = yaml.safe_load(MINIMAL_PROFILE.read_text(encoding="utf-8"))
            source = profile["sources"][0]
            source["id"] = "ads"
            source["platform"] = "meta"
            source["date_column"] = {"raw": "Day", "parse": "%Y-%m-%d"}
            source["observed_schema"] = [
                {"raw": "Day", "canonical": "date", "dtype": "date"},
                {"raw": "Account", "canonical": "account_id", "dtype": "text"},
                {"raw": "Ad", "canonical": "ad_id", "dtype": "text"},
                {"raw": "Spend", "canonical": "spend", "dtype": "number"},
            ]
            source["required_columns"] = ["date", "account_id", "ad_id", "spend"]
            profile["sources"] = [source]
            profile["source_plan"] = {
                "analysis_grain": "ad",
                "grains_provided": ["ad"],
                "joins": [],
                "additive_policy": {"forbid_cross_grain_sum": True, "additive_measures": ["spend"]},
            }
            profile["metrics"] = {"spend": {"role": "spend", "from": "ads.spend", "agg": "sum"}}
            profile["derived"] = {}
            profile["report_types"][0]["comparison"] = {"mode": "same_point_prev_day"}
            profile["report_types"][0]["aggregation"]["group_by"] = ["account_id"]
            profile["output"]["artifacts"][0]["grain"] = "ad"
            template_path = work / "template.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Compare"
            sheet.append(["日期", "账户ID", "花费"])
            book.save(template_path)
            profile["output"]["artifacts"].insert(0, {
                "kind": "xlsx_workbook",
                "write_mode": "append_to_sample",
                "target": {"mode": "append_to_sample", "sample": "template.xlsx"},
                "sheets": [{
                    "id": "comparison",
                    "match": "daily",
                    "target_sheet": "Compare",
                    "role": "comparison",
                    "row_key": {"field": "日期", "format": "date:yyyy-mm-dd"},
                    "placement": {"mode": "append_bottom", "same_key": "replace_in_place"},
                    "group_by": ["账户ID"],
                    "columns": [
                        {"header": "日期", "source": "$report_date", "format": "date:yyyy-mm-dd"},
                        {"header": "账户ID", "source": "group.account_id", "format": "text"},
                        {"header": "花费", "source": "metrics.spend", "format": "number:2"},
                    ],
                    "comparison_layout": {
                        "mode": "block_interleaved",
                        "block": {"rows": ["base", "current", "delta"], "delta_label": "环比", "delta_value": "raw_ratio", "null_symbol": "/"},
                    },
                }],
            })
            profile_path = work / "profile.yaml"
            profile_path.write_text(yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
            current_path = work / "current.csv"
            base_path = work / "base.csv"
            write_csv(current_path, [{"Day": "2026-10-25", "Account": "=1+1", "Ad": "ad-1", "Spend": 20}])
            write_csv(base_path, [{"Day": "2026-10-24", "Account": "=1+1", "Ad": "ad-1", "Spend": 10}])
            output_path = work / "output.xlsx"

            result = run_cli(
                SCRIPTS / "render_workbook.py",
                "--profile", profile_path,
                "--report-type", "daily",
                "--date", "2026-10-25",
                "--source-file", f"ads={current_path}",
                "--compare-source-file", f"ads={base_path}",
                "--out", output_path,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            output = load_workbook(output_path, data_only=False)
            values = [output["Compare"].cell(row=row, column=3).value for row in range(2, 5)]
            dates = [output["Compare"].cell(row=row, column=1).value for row in range(2, 4)]
            account_cell = output["Compare"].cell(row=2, column=2)
            output.close()
            self.assertEqual(values, [10.0, 20.0, 1.0])
            self.assertEqual(dates, ["2026-10-24", "2026-10-25"])
            self.assertEqual(account_cell.value, "'=1+1")
            self.assertEqual(account_cell.data_type, "s")

            broken = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            broken["output"]["artifacts"][0]["sheets"][0]["target_sheet"] = "Missing"
            broken_path = work / "broken-profile.yaml"
            broken_path.write_text(yaml.safe_dump(broken, allow_unicode=True, sort_keys=False), encoding="utf-8")
            broken_output = work / "broken-output.xlsx"
            failed = run_cli(
                SCRIPTS / "render_workbook.py",
                "--profile", broken_path,
                "--report-type", "daily",
                "--date", "2026-10-25",
                "--source-file", f"ads={current_path}",
                "--compare-source-file", f"ads={base_path}",
                "--out", broken_output,
            )
            self.assertNotEqual(failed.returncode, 0)
            self.assertFalse(broken_output.exists())

    def test_probe_xlsx_emits_columns_for_inference(self) -> None:
        test_tmp = ROOT / "runtime" / "test-tmp"
        test_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_tmp) as tmp:
            work = Path(tmp)
            source_path = work / "meta_ads.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(["Day", "Account ID", "Ad ID", "Spend"])
            sheet.append(["2026-10-25", "acct-1", "ad-1", 12])
            book.save(source_path)
            probe_path = work / "probe.json"

            result = run_cli(
                SCRIPTS / "probe_sources.py",
                "--source-file", source_path,
                "--out", probe_path,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            probe = json.loads(probe_path.read_text(encoding="utf-8"))
            source = probe["sources"][0]
            self.assertEqual([column["name"] for column in source["columns"]], ["Day", "Account ID", "Ad ID", "Spend"])
            self.assertTrue(source["grain_candidates"])

    def test_probe_recognizes_common_english_registration_headers(self) -> None:
        test_tmp = ROOT / "runtime" / "test-tmp"
        test_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_tmp) as tmp:
            work = Path(tmp)
            source_path = work / "adjust_events.csv"
            source_path.write_text(
                "Date,Ad ID,Registrations,Registrations Activity,Registrations LTV\n"
                "2026-10-25,ad-1,7,6,70\n",
                encoding="utf-8",
            )
            probe_path = work / "probe.json"

            result = run_cli(
                SCRIPTS / "probe_sources.py",
                "--source-file", source_path,
                "--out", probe_path,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            probe = json.loads(probe_path.read_text(encoding="utf-8"))
            roles = {
                column["name"]: column["role_candidates"]
                for column in probe["sources"][0]["columns"]
            }
            for name in ("Registrations", "Registrations Activity", "Registrations LTV"):
                self.assertEqual(roles[name], ["registration"])

    def test_xlsx_source_uses_profile_layout_and_generates_report(self) -> None:
        test_tmp = ROOT / "runtime" / "test-tmp"
        test_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_tmp) as tmp:
            work = Path(tmp)
            profile = yaml.safe_load(MINIMAL_PROFILE.read_text(encoding="utf-8"))
            source = profile["sources"][0]
            source["id"] = "ads"
            source["platform"] = "meta"
            source["layout"] = {"header_row": 1, "sheet": "Export", "encoding": "utf-8"}
            source["date_column"] = {"raw": "Day", "parse": "%Y-%m-%d"}
            source["observed_schema"] = [
                {"raw": "Day", "canonical": "date", "dtype": "date"},
                {"raw": "Account", "canonical": "account_id", "dtype": "text"},
                {"raw": "Ad", "canonical": "ad_id", "dtype": "text"},
                {"raw": "Spend", "canonical": "spend", "dtype": "number"},
            ]
            source["required_columns"] = ["date", "account_id", "ad_id", "spend"]
            profile["sources"] = [source]
            profile["source_plan"] = {
                "analysis_grain": "ad",
                "grains_provided": ["ad"],
                "joins": [],
                "additive_policy": {"forbid_cross_grain_sum": True, "additive_measures": ["spend"]},
            }
            profile["metrics"] = {"spend": {"role": "spend", "from": "ads.spend", "agg": "sum"}}
            profile["derived"] = {}
            profile["output"]["artifacts"][0]["grain"] = "ad"
            profile["output"]["artifacts"][1]["group_by"] = ["账户ID"]
            profile_path = work / "profile.yaml"
            profile_path.write_text(yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8")

            source_path = work / "ads.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "Export"
            sheet.append(["generated export"])
            sheet.append(["Day", "Account", "Ad", "Spend"])
            sheet.append(["2026-10-25", "acct-1", "ad-1", 12])
            book.save(source_path)

            out_dir = work / "out"
            result = run_cli(
                SCRIPTS / "generate_report.py",
                "--profile", profile_path,
                "--report-type", "daily",
                "--date", "2026-10-25",
                "--source-file", f"ads={source_path}",
                "--out-dir", out_dir,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            with (out_dir / "过程表_20261025.csv").open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["花费"], "12")

    def test_unknown_report_type_is_rejected_before_reading_sources(self) -> None:
        test_tmp = ROOT / "runtime" / "test-tmp"
        test_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_tmp) as tmp:
            profile_path = Path(tmp) / "profile.yaml"
            profile_path.write_text(MINIMAL_PROFILE.read_text(encoding="utf-8"), encoding="utf-8")
            result = run_cli(
                SCRIPTS / "generate_report.py",
                "--profile", profile_path,
                "--report-type", "weekly_missing",
                "--date", "2026-10-25",
                "--out-dir", Path(tmp) / "out",
            )
            self.assertEqual(result.returncode, 2)
            self.assertIn("不存在的报表类型", result.stderr)

    def test_source_plan_drives_generic_join_dedupe_and_unmatched_status(self) -> None:
        test_tmp = ROOT / "runtime" / "test-tmp"
        test_tmp.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=test_tmp) as tmp:
            work = Path(tmp)
            profile = yaml.safe_load(MINIMAL_PROFILE.read_text(encoding="utf-8"))
            profile["sources"] = [
                {
                    "id": "ads",
                    "platform": "meta",
                    "role": "ad_platform",
                    "grain": "ad",
                    "discover": {"dir": "inbox", "glob": "ads.csv", "pick": {"strategy": "explicit"}},
                    "layout": {"header_row": 0, "encoding": "utf-8"},
                    "date_column": {"raw": "Day", "parse": "%Y-%m-%d"},
                    "observed_schema": [
                        {"raw": "Day", "canonical": "date", "dtype": "date"},
                        {"raw": "Account", "canonical": "account_id", "dtype": "text"},
                        {"raw": "Ad", "canonical": "ad_id", "dtype": "text"},
                        {"raw": "Spend", "canonical": "spend", "dtype": "number"},
                    ],
                    "required_columns": ["date", "account_id", "ad_id", "spend"],
                    "snapshot": {"is_snapshot": False},
                },
                {
                    "id": "attrib",
                    "platform": "adjust",
                    "role": "attribution",
                    "grain": "ad",
                    "discover": {"dir": "inbox", "glob": "attrib.csv", "pick": {"strategy": "explicit"}},
                    "layout": {"header_row": 0, "encoding": "utf-8"},
                    "date_column": {"raw": "Day", "parse": "%Y-%m-%d"},
                    "observed_schema": [
                        {"raw": "Day", "canonical": "date", "dtype": "date"},
                        {"raw": "Account", "canonical": "account_id", "dtype": "text"},
                        {"raw": "Ad", "canonical": "ad_id", "dtype": "text"},
                        {"raw": "Snapshot", "canonical": "snapshot_date", "dtype": "date"},
                        {"raw": "Registrations", "canonical": "registration", "dtype": "number"},
                    ],
                    "required_columns": ["date", "account_id", "ad_id", "registration"],
                    "snapshot": {"is_snapshot": True, "snapshot_field": "snapshot_date"},
                },
            ]
            profile["source_plan"] = {
                "analysis_grain": "ad",
                "grains_provided": ["ad"],
                "joins": [{
                    "id": "ads_x_attrib",
                    "left": "ads",
                    "right": "attrib",
                    "type": "left",
                    "keys": ["date", "account_id", "ad_id"],
                    "enforce_grain": "ad",
                }],
                "additive_policy": {
                    "forbid_cross_grain_sum": True,
                    "additive_measures": ["spend", "registration"],
                },
                "dedupe": [{
                    "source": "attrib",
                    "key": ["date", "account_id", "ad_id"],
                    "keep": "latest_snapshot",
                }],
                "unmatched_policy": {"keep_left": True, "report": True, "label": "未匹配"},
            }
            profile["metrics"] = {
                "spend": {"role": "spend", "from": "ads.spend", "agg": "sum"},
                "registration": {"role": "registration", "from": "attrib.registration", "agg": "sum"},
            }
            profile["derived"] = {}
            profile["rules"] = {"join_coverage_threshold": 0.5}
            profile["output"]["artifacts"][0]["grain"] = "ad"
            profile["output"]["artifacts"][1]["group_by"] = ["账户ID"]
            profile["output"]["artifacts"][1].pop("rows_filter", None)
            profile["validation"]["checks"] = [
                {"id": "total_reconcile", "rule": "totals", "level": "blocking"}
            ]

            profile_path = work / "profile.yaml"
            profile_path.write_text(yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
            ads_path = work / "ads.csv"
            attrib_path = work / "attrib.csv"
            write_csv(ads_path, [
                {"Day": "2026-10-25", "Account": "acct-1", "Ad": "ad-1", "Spend": 10},
                {"Day": "2026-10-25", "Account": "acct-1", "Ad": "ad-2", "Spend": 20},
            ])
            write_csv(attrib_path, [
                {"Day": "2026-10-25", "Account": "acct-1", "Ad": "ad-1", "Snapshot": "2026-10-25", "Registrations": 2},
                {"Day": "2026-10-25", "Account": "acct-1", "Ad": "ad-1", "Snapshot": "2026-10-26", "Registrations": 5},
            ])
            out_dir = work / "out"

            result = run_cli(
                SCRIPTS / "generate_report.py",
                "--profile", profile_path,
                "--report-type", "daily",
                "--date", "2026-10-25",
                "--source-file", f"ads={ads_path}",
                "--source-file", f"attrib={attrib_path}",
                "--out-dir", out_dir,
            )

            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            with (out_dir / "过程表_20261025.csv").open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual([row["registration"] for row in rows], ["5", "0"])
            self.assertEqual([row["匹配状态"] for row in rows], ["已匹配", "未匹配"])

            strict_profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            strict_profile["rules"]["join_coverage_threshold"] = 0.75
            strict_path = work / "strict-profile.yaml"
            strict_path.write_text(yaml.safe_dump(strict_profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
            strict_out = work / "strict-out"
            low_coverage = run_cli(
                SCRIPTS / "generate_report.py",
                "--profile", strict_path,
                "--report-type", "daily",
                "--date", "2026-10-25",
                "--source-file", f"ads={ads_path}",
                "--source-file", f"attrib={attrib_path}",
                "--out-dir", strict_out,
            )
            self.assertNotEqual(low_coverage.returncode, 0)
            self.assertIn("关联覆盖率", low_coverage.stderr)
            self.assertFalse((strict_out / "过程表_20261025.csv").exists())

            maturity_profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            maturity_profile["sources"][1]["snapshot"]["maturity"] = {"registration": 7}
            maturity_path = work / "maturity-profile.yaml"
            maturity_path.write_text(yaml.safe_dump(maturity_profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
            maturity_out = work / "maturity-out"
            immature = run_cli(
                SCRIPTS / "generate_report.py",
                "--profile", maturity_path,
                "--report-type", "daily",
                "--date", "2026-10-25",
                "--source-file", f"ads={ads_path}",
                "--source-file", f"attrib={attrib_path}",
                "--out-dir", maturity_out,
            )
            self.assertNotEqual(immature.returncode, 0)
            self.assertIn("快照指标未成熟", immature.stderr)
            self.assertFalse((maturity_out / "过程表_20261025.csv").exists())

            inner_profile = yaml.safe_load(profile_path.read_text(encoding="utf-8"))
            inner_profile["source_plan"]["joins"][0]["type"] = "inner"
            inner_path = work / "inner-profile.yaml"
            inner_path.write_text(yaml.safe_dump(inner_profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
            inner_out = work / "inner-out"
            inner = run_cli(
                SCRIPTS / "generate_report.py",
                "--profile", inner_path,
                "--report-type", "daily",
                "--date", "2026-10-25",
                "--source-file", f"ads={ads_path}",
                "--source-file", f"attrib={attrib_path}",
                "--out-dir", inner_out,
            )
            self.assertEqual(inner.returncode, 0, inner.stdout + inner.stderr)
            with (inner_out / "过程表_20261025.csv").open(encoding="utf-8-sig", newline="") as handle:
                inner_rows = list(csv.DictReader(handle))
            self.assertEqual(len(inner_rows), 1)

            missing_out = work / "missing-out"
            missing = run_cli(
                SCRIPTS / "generate_report.py",
                "--profile", profile_path,
                "--report-type", "daily",
                "--date", "2026-10-25",
                "--source-file", f"ads={ads_path}",
                "--out-dir", missing_out,
            )
            self.assertNotEqual(missing.returncode, 0)
            self.assertIn("缺少必需数据源", missing.stderr)
            self.assertFalse((missing_out / "过程表_20261025.csv").exists())

            invalid_attrib = work / "invalid-attrib.csv"
            write_csv(invalid_attrib, [{
                "Day": "2026-10-25", "Account": "acct-1", "Ad": "ad-1",
                "Snapshot": "2026-10-26", "Registrations": "not-a-number",
            }])
            invalid_out = work / "invalid-out"
            invalid = run_cli(
                SCRIPTS / "generate_report.py",
                "--profile", profile_path,
                "--report-type", "daily",
                "--date", "2026-10-25",
                "--source-file", f"ads={ads_path}",
                "--source-file", f"attrib={invalid_attrib}",
                "--out-dir", invalid_out,
            )
            self.assertNotEqual(invalid.returncode, 0)
            self.assertIn("无法解析为数字", invalid.stderr)
            self.assertFalse((invalid_out / "过程表_20261025.csv").exists())


if __name__ == "__main__":
    unittest.main()
