#!/usr/bin/env python3
"""Create the three-sheet standard workbook used when no target template exists.

This command only creates workbook structure. It never edits the profile, source
files, or an existing workbook. A period must be explicitly confirmed before any
workbook can be created.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from cli_compat import configure_console

configure_console()

try:
    import yaml
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    print(f"缺少依赖：{exc.name}，请先执行 pip install -r requirements.txt", file=sys.stderr)
    sys.exit(2)

sys.path.insert(0, str(Path(__file__).resolve().parent))
import validate_profile as vp  # noqa: E402


SHEET_NAMES = ["管理摘要", "明细数据", "趋势环比"]
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRAY = "666666"


def pending_p0(profile: dict[str, Any]) -> list[dict[str, Any]]:
    return [
        item for item in (profile.get("ambiguities") or [])
        if item.get("level") == "P0" and item.get("status") == "pending"
    ]


def metric_names(profile: dict[str, Any]) -> list[str]:
    requested = list(((profile.get("report_intent") or {}).get("primary_kpis") or []))
    available = profile.get("metrics") or {}
    names = [name for name in requested if name in available]
    return (names or list(available))[:8]


def display_name(name: str) -> str:
    labels = {
        "spend": "花费", "impressions": "展示", "clicks": "点击", "install": "安装",
        "registration": "注册", "kyc": "KYC", "first_deposit": "首充", "trade": "交易",
        "revenue": "收入",
    }
    return labels.get(name, name)


def set_title(ws, title: str, *, draft: bool) -> None:
    ws.merge_cells("A1:H1")
    cell = ws["A1"]
    cell.value = "DRAFT｜待确认｜非正式交付" if draft else title
    cell.font = Font(size=15, bold=True, color=BLUE if not draft else "9C6500")
    cell.fill = PatternFill("solid", fgColor=YELLOW if draft else WHITE)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26


def style_header(ws, row: int, columns: int) -> None:
    thin = Side(style="thin", color="B7C9D6")
    for column in range(1, columns + 1):
        cell = ws.cell(row=row, column=column)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)


def style_columns(ws, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def create_summary(ws, profile: dict[str, Any], *, draft: bool) -> None:
    intent = profile["report_intent"]
    set_title(ws, "管理摘要", draft=draft)
    ws["A2"] = (profile.get("product") or {}).get("name", "")
    ws["A2"].font = Font(size=18, bold=True, color=BLUE)
    ws["A3"] = "业务目标"
    ws["B3"] = intent.get("goal_description") or intent.get("primary_goal") or ""
    ws["A4"] = "报告周期"
    period = intent.get("period") or {}
    ws["B4"] = period.get("type") or ""
    if period.get("start") or period.get("end"):
        ws["C4"] = f"{period.get('start', '')} ~ {period.get('end', '')}"
    ws["A5"] = "AI 数据结论"
    ws["B5"] = "待正式数据生成后，仅基于已验证数字填写。"
    ws["B5"].font = Font(italic=True, color=GRAY)

    metrics = metric_names(profile)
    headers = ["周期", "对比", *[display_name(name) for name in metrics]]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=7, column=column, value=header)
    style_header(ws, 7, len(headers))
    for offset, (period_label, comparison_label) in enumerate([
        ("上一等长周期", "上期"), ("本期", "本期"), ("绝对变化", "绝对变化"), ("变化率", "环比")
    ], start=8):
        ws.cell(row=offset, column=1, value=period_label)
        ws.cell(row=offset, column=2, value=comparison_label)
        for column in range(3, len(headers) + 1):
            ws.cell(row=offset, column=column, value="N/A" if offset == 11 else None)
            ws.cell(row=offset, column=column).number_format = "0.00%" if offset == 11 else "#,##0.00"
    ws.freeze_panes = "A8"
    ws.auto_filter.ref = f"A7:{get_column_letter(len(headers))}8"
    style_columns(ws, [18, *([15] * len(metrics))])

    chart = LineChart()
    chart.title = "核心指标趋势"
    chart.y_axis.title = "指标值"
    chart.x_axis.title = "周期"
    if metrics:
        chart.add_data(Reference(ws, min_col=3, max_col=len(headers), min_row=7, max_row=9), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=8, max_row=9))
    chart.height = 7
    chart.width = 15
    ws.add_chart(chart, "A11")


def create_detail(ws, profile: dict[str, Any], *, draft: bool) -> None:
    set_title(ws, "明细数据", draft=draft)
    intent = profile["report_intent"]
    dimension_labels = {
        "date": "日期", "account_id": "账户ID", "country_code": "国家代码",
        "campaign_id": "Campaign ID", "adgroup_id": "Ad Group ID", "ad_id": "Ad ID",
    }
    dimensions = [name for name in (intent.get("dimensions") or []) if name in dimension_labels]
    headers = [dimension_labels[name] for name in dimensions]
    if "日期" not in headers:
        headers.insert(0, "日期")
    headers.extend(display_name(name) for name in metric_names(profile))
    for column, header in enumerate(headers, start=1):
        ws.cell(row=3, column=column, value=header)
    style_header(ws, 3, len(headers))
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}3"
    style_columns(ws, [15 if header == "日期" else 18 for header in headers])


def create_trend(ws, profile: dict[str, Any], *, draft: bool) -> None:
    set_title(ws, "趋势环比", draft=draft)
    headers = ["周期", "对比", *[display_name(name) for name in metric_names(profile)]]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=3, column=column, value=header)
    style_header(ws, 3, len(headers))
    ws["A4"] = "上一等长周期"
    ws["B4"] = "上期"
    ws["A5"] = "本期"
    ws["B5"] = "本期"
    ws["A6"] = "绝对变化"
    ws["B6"] = "绝对变化"
    ws["A7"] = "变化率"
    ws["B7"] = "环比"
    for column in range(3, len(headers) + 1):
        ws.cell(row=7, column=column, value="N/A")
        ws.cell(row=7, column=column).number_format = "0.00%"
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}7"
    style_columns(ws, [18, 12, *([15] * max(0, len(headers) - 2))])


def output_path(requested: Path, draft: bool) -> Path:
    if not draft or requested.stem.endswith("_DRAFT"):
        return requested
    return requested.with_name(f"{requested.stem}_DRAFT{requested.suffix}")


def main() -> int:
    parser = argparse.ArgumentParser(description="按已确认 profile 创建标准三 Sheet 模板")
    parser.add_argument("--profile", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--draft", action="store_true", help="生成带黄色待确认标识的草稿")
    parser.add_argument("--force", action="store_true", help="显式覆盖已存在的目标文件")
    args = parser.parse_args()

    if not args.profile.exists():
        print(f"找不到 profile：{args.profile}", file=sys.stderr)
        return 2
    try:
        profile = yaml.safe_load(args.profile.read_text(encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError) as exc:
        print(f"无法读取 profile：{exc}", file=sys.stderr)
        return 2
    intent = profile.get("report_intent") or {}
    period = intent.get("period") or {}
    if not period.get("confirmed") or period.get("type") not in {"daily", "weekly", "monthly", "custom"}:
        recommendation = period.get("recommendation") or {}
        suggested = recommendation.get("suggested_type") or "daily / weekly / monthly / custom"
        print(
            f"时间周期尚未由用户确认，已阻断生成。建议选择：{suggested}；"
            "请在 report_intent.period 写入 type 并设置 confirmed: true。",
            file=sys.stderr,
        )
        return 1
    structural_failures = [issue for issue in vp.schema_issues(profile) if issue.level == "FAIL"]
    if structural_failures:
        print("profile 结构校验失败，未生成任何工作簿：", file=sys.stderr)
        for issue in structural_failures:
            print(f"  {issue}", file=sys.stderr)
        return 1
    unresolved = pending_p0(profile)
    if unresolved and not args.draft:
        print("存在未解决 P0，只能使用 --draft 生成预览：", file=sys.stderr)
        for item in unresolved:
            print(f"  {item.get('id')}: {item.get('question')}", file=sys.stderr)
        return 1
    if not args.draft and not intent.get("layout_confirmed", False):
        print("版式尚未确认，只能先使用 --draft 生成预览。", file=sys.stderr)
        return 1
    allowed_draft_failures = {"UNRESOLVED_P0", "UNCONFIRMED_LAYOUT"}
    semantic_failures = [
        issue for issue in vp.semantic_issues(profile)
        if issue.level == "FAIL" and (not args.draft or issue.code not in allowed_draft_failures)
    ]
    if semantic_failures:
        print("profile 语义校验失败，未生成任何工作簿：", file=sys.stderr)
        for issue in semantic_failures:
            print(f"  {issue}", file=sys.stderr)
        return 1

    target = output_path(args.out, args.draft)
    if target.exists() and not args.force:
        print(f"目标文件已存在，未覆盖：{target}（如需覆盖请显式传入 --force）", file=sys.stderr)
        return 1

    workbook = Workbook()
    summary = workbook.active
    summary.title = SHEET_NAMES[0]
    detail = workbook.create_sheet(SHEET_NAMES[1])
    trend = workbook.create_sheet(SHEET_NAMES[2])
    create_summary(summary, profile, draft=args.draft)
    create_detail(detail, profile, draft=args.draft)
    create_trend(trend, profile, draft=args.draft)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    target.parent.mkdir(parents=True, exist_ok=True)
    staged = target.with_name(f".{target.stem}.tmp{target.suffix}")
    workbook.save(staged)
    workbook.close()
    staged.replace(target)
    print(f"标准{'草稿' if args.draft else '模板'}已生成：{target}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
