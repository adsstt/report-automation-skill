from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parent
HEADERS = [
    "Day", "Account ID", "Account Name", "Campaign ID", "Campaign Name",
    "Ad Set ID", "Ad Set Name", "Ad ID", "Ad Name", "Spend",
    "Impressions", "Clicks", "Installs",
]


def meta_rows(day: str, factor: float) -> list[list[object]]:
    definitions = [
        ("US", "us-1", 100, 10000, 500, 100),
        ("US", "us-2", 50, 5000, 250, 50),
        ("GB", "gb-1", 80, 8000, 400, 80),
        ("GB", "gb-2", 70, 7000, 350, 70),
    ]
    return [[
        day, "nova-meta-001", "NovaFit Growth", f"cmp-{country}", f"NF_{country}_REG",
        f"set-{country}", f"Prospecting {country}", f"ad-{token}", f"Creative {token}",
        round(spend * factor, 2), int(impressions * factor), int(clicks * factor), int(installs * factor),
    ] for country, token, spend, impressions, clicks, installs in definitions]


def write_meta(day: str, factor: float) -> None:
    token = day.replace("-", "")
    path = ROOT / "sources" / day / f"meta_ads_{token}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Export"
    sheet.append([f"NovaFit Meta export generated for {day}"])
    sheet.append(HEADERS)
    for row in meta_rows(day, factor):
        sheet.append(row)
    book.save(path)


def write_adjust(day: str, current: bool) -> None:
    token = day.replace("-", "")
    path = ROOT / "sources" / day / f"adjust_attribution_{token}.csv"
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = (
        [
            [day, "nova-meta-001", "ad-us-1", day, 8, 5, 2, 200],
            [day, "nova-meta-001", "ad-us-1", "2026-10-26", 10, 6, 3, 300],
            [day, "nova-meta-001", "ad-us-2", "2026-10-26", 5, 3, 1, 100],
            [day, "nova-meta-001", "ad-gb-1", "2026-10-26", 8, 5, 2, 200],
            [day, "nova-meta-001", "ad-gb-2", "2026-10-26", 7, 4, 2, 200],
        ] if current else [
            [day, "nova-meta-001", "ad-us-1", day, 8, 5, 2, 220],
            [day, "nova-meta-001", "ad-us-2", day, 4, 2, 1, 80],
            [day, "nova-meta-001", "ad-gb-1", day, 6, 4, 1, 120],
            [day, "nova-meta-001", "ad-gb-2", day, 6, 3, 2, 180],
        ]
    )
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Date", "Account ID", "Ad ID", "Snapshot Date", "signup_success", "kyc_successed", "first_deposited", "Revenue USD"])
        writer.writerows(rows)


def write_template() -> None:
    path = ROOT / "samples" / "target_template.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    history = book.active
    history.title = "Daily History"
    history.append(["日期", "账户名称", "账户ID", "国家名称", "国家代码", "事件", "花费", "注册", "KYC", "FTD", "收入", "注册成本"])
    comparison = book.create_sheet("Daily Comparison")
    comparison.append(["日期", "国家", "花费", "注册", "KYC", "FTD", "收入"])
    for sheet in (history, comparison):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="305496")
        sheet.freeze_panes = "A2"
    book.save(path)


if __name__ == "__main__":
    write_meta("2026-10-24", 0.8)
    write_meta("2026-10-25", 1.0)
    write_adjust("2026-10-24", current=False)
    write_adjust("2026-10-25", current=True)
    write_template()
    print(f"Fresh-product fixtures created under {ROOT}")
