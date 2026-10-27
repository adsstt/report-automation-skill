#!/usr/bin/env python3
"""源文件与目标格式探针（probe）

只做"看"，不做判断：把源文件和目标报表的客观结构全部摸出来，供 AI 反推规则、
供 infer_mappings.py 生成候选。这一步不产出任何规则，也不修改任何文件。

输出一份 JSON 观测报告 + 一份人类可读摘要。

用法：
    python scripts/probe_sources.py --source-dir <dir> [--source-dir <dir> ...] --out probe.json
    python scripts/probe_sources.py --source-dir <dir> --target <目标报表.xlsx|.csv> --out probe.json

设计要点：
- CSV 编码、分隔符、表头行都不假设，靠探测。
- 表头行不固定在第 0 行：在前 10 行里找"非空单元格最多且在后续行中被稳定复用"的那一行。
- 日期列通过"值能解析成日期的比例"识别，不靠列名猜。
- 粒度（campaign/adgroup/ad/creative）通过 ID 类列名的模式匹配给出**候选**，标注置信度，
  最终由人确认——探针不下结论。
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d", "%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S")
ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "cp936", "latin-1")

# 列名 → 语义角色候选（只给候选，不下结论）
COLUMN_HINTS: dict[str, list[str]] = {
    "account_id": [r"^账户\s*id$", r"^账号\s*id$", r"^account\s*id$"],
    "account_name": [r"账户名称", r"账号名称", r"^account\s*name$"],
    "campaign_id": [r"推广系列\s*id", r"系列\s*id", r"^campaign\s*id$"],
    "campaign_name": [r"系列名称", r"^campaign$", r"^campaign\s*name$"],
    "adgroup_id": [r"广告组\s*id", r"^adset\s*id$", r"^ad\s*group\s*id$"],
    "adgroup_name": [r"广告组名称", r"^adset\s*name$", r"^ad\s*group\s*name$"],
    "ad_id": [r"^广告\s*id$", r"^ad\s*id$"],
    "ad_name": [r"广告名称", r"^ad\s*name$"],
    "creative_id": [r"创意素材\s*id", r"^creative\s*id$"],
    "creative_name": [r"创意素材名称", r"^creative\s*name$"],
    "spend": [r"^消耗$", r"^花费$", r"^spend$", r"^cost$"],
    "impressions": [r"^展示量$", r"^曝光", r"^impressions$"],
    "clicks": [r"点击量", r"^clicks?$"],
    "install": [r"安装", r"^installs?"],
    "registration": [r"signup_success", r"注册"],
    "kyc": [r"kyc_successed"],
    "first_deposit": [r"first_deposited"],
    "trade": [r"first_traded", r"trade_all"],
    "revenue": [r"^revenue", r"收入"],
    "date": [r"^按天$", r"^日期$", r"^date$", r"^day$"],
}

GRAIN_ID_KEYS = {
    "campaign": ["campaign_id"],
    "adgroup": ["adgroup_id"],
    "ad": ["ad_id"],
    "creative": ["creative_id"],
}


def sniff_encoding(path: Path) -> tuple[str, str]:
    """返回 (encoding, text)。CSV 编码不假设。"""
    raw = path.read_bytes()
    for enc in ENCODINGS:
        try:
            return enc, raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return "latin-1", raw.decode("latin-1", errors="replace")


def sniff_delimiter(text: str) -> str:
    head = "\n".join(text.splitlines()[:5])
    counts = {d: head.count(d) for d in [",", "\t", ";", "|"]}
    return max(counts, key=counts.get) if max(counts.values()) > 0 else ","


def detect_header_row(rows: list[list[str]], max_scan: int = 10) -> int:
    """在前 N 行里找表头行：非空单元格最多、且后续行也普遍非空的那一行。"""
    best_idx, best_score = 0, -1
    for idx, row in enumerate(rows[:max_scan]):
        non_empty = sum(1 for c in row if str(c).strip())
        if non_empty == 0:
            continue
        # 后续 5 行的平均非空数，衡量这一行是否像表头
        following = rows[idx + 1: idx + 6]
        follow_filled = (
            sum(sum(1 for c in r if str(c).strip()) for r in following) / max(len(following), 1)
            if following
            else 0
        )
        score = non_empty + follow_filled * 0.1
        if score > best_score:
            best_idx, best_score = idx, score
    return best_idx


def parse_date(value: str) -> datetime | None:
    v = str(value).strip()
    if not v:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(v, fmt)
        except ValueError:
            continue
    return None


def role_for(column: str) -> list[str]:
    hits = []
    low = column.strip().lower()
    for role, patterns in COLUMN_HINTS.items():
        if any(re.search(p, low, flags=re.IGNORECASE) for p in patterns):
            hits.append(role)
    return hits


def probe_csv(path: Path) -> dict[str, Any]:
    enc, text = sniff_encoding(path)
    delim = sniff_delimiter(text)
    raw_rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    if not raw_rows:
        return {"path": str(path), "error": "空文件"}
    header_idx = detect_header_row(raw_rows)
    header = [str(c).strip() for c in raw_rows[header_idx]]
    data_rows = raw_rows[header_idx + 1:]

    col_stats: list[dict[str, Any]] = []
    for ci, name in enumerate(header):
        values = [r[ci].strip() for r in data_rows if ci < len(r)]
        non_empty = [v for v in values if v not in ("", "-", "--")]
        numeric = 0
        for v in non_empty[:500]:
            try:
                float(v.replace(",", "").replace("%", "").replace("$", ""))
                numeric += 1
            except ValueError:
                pass
        dated = sum(1 for v in non_empty[:500] if parse_date(v))
        distinct = len(set(non_empty))
        col_stats.append(
            {
                "index": ci,
                "name": name,
                "non_empty": len(non_empty),
                "numeric_ratio": round(numeric / max(len(non_empty[:500]), 1), 3),
                "date_ratio": round(dated / max(len(non_empty[:500]), 1), 3),
                "distinct": distinct,
                "sample": non_empty[:3],
                "role_candidates": role_for(name),
            }
        )

    total_rows = [
        i for i, r in enumerate(data_rows)
        if any(str(c).strip() in {"总计", "合计", "Total", "Grand Total"} for c in r[:3])
    ]

    return {
        "path": str(path),
        "format": "csv",
        "encoding": enc,
        "delimiter": delim,
        "header_row": header_idx,
        "row_count": len(data_rows),
        "columns": col_stats,
        "total_like_rows": total_rows[:5],
        "total_like_row_count": len(total_rows),
    }


def probe_xlsx(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"path": str(path), "error": "缺少 openpyxl，无法读 xlsx"}

    wb = load_workbook(filename=path, data_only=True, read_only=True)
    sheets = []
    primary: dict[str, Any] = {}
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append(["" if c is None else str(c) for c in row])
            if i >= 500:
                break
        header_idx = detect_header_row(rows) if rows else 0
        header = rows[header_idx] if rows else []
        data_rows = rows[header_idx + 1:]
        columns: list[dict[str, Any]] = []
        for ci, name in enumerate(header):
            if not str(name).strip():
                continue
            values = [row[ci].strip() for row in data_rows if ci < len(row)]
            non_empty = [value for value in values if value not in ("", "-", "--")]
            numeric = 0
            for value in non_empty:
                try:
                    float(value.replace(",", "").replace("%", "").replace("$", ""))
                    numeric += 1
                except ValueError:
                    pass
            dated = sum(1 for value in non_empty if parse_date(value))
            columns.append({
                "index": ci,
                "name": str(name).strip(),
                "non_empty": len(non_empty),
                "numeric_ratio": round(numeric / max(len(non_empty), 1), 3),
                "date_ratio": round(dated / max(len(non_empty), 1), 3),
                "distinct": len(set(non_empty)),
                "sample": non_empty[:3],
                "role_candidates": role_for(str(name)),
            })
        sheets.append(
            {
                "name": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "header_row_guess": header_idx,
                "header": [c for c in header if str(c).strip()][:60],
                "first_rows": rows[: min(6, len(rows))],
            }
        )
        if not primary and columns:
            dates = []
            dated = next((column for column in columns if column["date_ratio"] > 0.8), None)
            if dated:
                index = dated["index"]
                dates = sorted({parsed for parsed in (parse_date(row[index]) for row in data_rows if index < len(row)) if parsed})
            primary = {
                "sheet": ws.title,
                "header_row": header_idx,
                "row_count": len(data_rows),
                "columns": columns,
                "date_coverage": ({
                    "column": dated["name"],
                    "min": dates[0].strftime("%Y-%m-%d"),
                    "max": dates[-1].strftime("%Y-%m-%d"),
                    "distinct_days": len({item.date() for item in dates}),
                } if dated and dates else {}),
            }
    wb.close()
    return {"path": str(path), "format": "xlsx", "sheets": sheets, **primary}


def detect_header_row_typed(rows: list[list[Any]], max_scan: int = 10) -> int:
    """表头行检测（面向目标报表）：取"非空且以文本为主"的行，而不是最宽的行。

    目标工作簿的数据行往往比表头更宽（表头只覆盖实际用到的列），
    用"最宽"会把数据行误判成表头。
    """
    best_idx, best_score = 0, -1.0
    for idx, row in enumerate(rows[:max_scan]):
        cells = [c for c in row if str(c).strip()]
        if not cells:
            continue
        # 含公式的行是"合计/汇总"行，不是表头
        if any(str(c).strip().startswith("=") for c in cells):
            continue
        numeric = 0
        for c in cells:
            try:
                float(str(c).replace(",", "").replace("%", "").replace("$", ""))
                numeric += 1
            except ValueError:
                pass
        text_ratio = 1 - numeric / len(cells)
        # 表头通常：文本占比高、列数不少
        score = text_ratio * 2 + min(len(cells), 30) / 30
        if score > best_score:
            best_idx, best_score = idx, score
    return best_idx


def probe_target_workbook(path: Path) -> dict[str, Any]:
    """目标报表探针：除了表头，还要看公式、合并单元格、条件格式，用于判断 sheet 角色。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"path": str(path), "error": "缺少 openpyxl"}

    wb = load_workbook(filename=path, data_only=False)
    info = []
    for ws in wb.worksheets:
        formulas = 0
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                if formulas > 200:
                    break
            if formulas > 200:
                break
        header = []
        rows_scan: list[list[Any]] = []
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True)):
            rows_scan.append(list(row))
        if rows_scan:
            hdr_idx = detect_header_row_typed(rows_scan)
            header = [str(c) for c in rows_scan[hdr_idx] if c not in (None, "")]
        info.append(
            {
                "name": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "merged_ranges": len(ws.merged_cells.ranges),
                "conditional_formatting_rules": sum(
                    len(r.rules) for r in getattr(ws, "conditional_formatting", [])
                ),
                "formula_cells_sampled": formulas,
                "header_guess": header[:60],
            }
        )
    wb.close()
    return {"path": str(path), "format": "xlsx", "sheets": info}


def grain_candidates(columns: list[dict[str, Any]], filename: str) -> list[dict[str, Any]]:
    present = set()
    for col in columns:
        present.update(col.get("role_candidates") or [])
    out = []
    for grain, keys in GRAIN_ID_KEYS.items():
        if all(k in present for k in keys):
            out.append({"grain": grain, "evidence": keys, "confidence": "high"})
        elif any(k in present for k in keys):
            out.append({"grain": grain, "evidence": keys, "confidence": "low"})
    low_name = filename.lower()
    for token, grain in (("广告id", "ad"), ("ad_id", "ad"), ("广告系列", "campaign"), ("campaign", "campaign"), ("广告组", "adgroup"), ("adset", "adgroup")):
        if token in low_name and not any(o["grain"] == grain for o in out):
            out.append({"grain": grain, "evidence": [f"文件名含 {token}"], "confidence": "medium"})
    return out


def date_coverage(columns: list[dict[str, Any]], path: Path, encoding: str, delim: str, header_row: int) -> dict[str, Any]:
    dated_cols = [c for c in columns if (c.get("date_ratio") or 0) > 0.8]
    if not dated_cols:
        return {}
    enc, text = sniff_encoding(path)
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))[header_row + 1:]
    idx = dated_cols[0]["index"]
    dates = sorted({d for d in (parse_date(r[idx]) for r in rows if idx < len(r)) if d})
    if not dates:
        return {}
    return {
        "column": dated_cols[0]["name"],
        "min": dates[0].strftime("%Y-%m-%d"),
        "max": dates[-1].strftime("%Y-%m-%d"),
        "distinct_days": len({d.date() for d in dates}),
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="探源：读取源文件与目标报表的结构")
    ap.add_argument("--source-dir", action="append", default=[], help="源文件目录，可多次指定")
    ap.add_argument("--source-file", action="append", default=[], help="单个源文件，可多次指定")
    ap.add_argument("--target", help="目标报表样本（xlsx 或 csv）")
    ap.add_argument("--glob", default="*", help="在 source-dir 中筛选文件名的通配符")
    ap.add_argument("--limit", type=int, default=5, help="每个目录最多探几个文件")
    ap.add_argument("--out", help="把 JSON 观测报告写到这个路径")
    args = ap.parse_args()

    files: list[Path] = [Path(p) for p in args.source_file]
    for d in args.source_dir:
        base = Path(d)
        if not base.exists():
            print(f"警告：目录不存在 {d}", file=sys.stderr)
            continue
        files.extend(sorted(p for p in base.glob(args.glob) if p.is_file())[-args.limit:])

    if not files and not args.target:
        print("至少需要 --source-dir / --source-file / --target 之一", file=sys.stderr)
        return 2

    report: dict[str, Any] = {"generated_at": datetime.now().isoformat(timespec="seconds"), "sources": []}

    print(f"探测 {len(files)} 个源文件")
    for path in files:
        try:
            if path.suffix.lower() in {".xlsx", ".xlsm"}:
                info = probe_xlsx(path)
                columns = info.get("columns", [])
            else:
                info = probe_csv(path)
                columns = info.get("columns", [])
                if columns:
                    info["date_coverage"] = date_coverage(columns, path, info["encoding"], info["delimiter"], info["header_row"])
            info["grain_candidates"] = grain_candidates(columns, path.name)
        except Exception as exc:  # noqa: BLE001
            info = {"path": str(path), "error": f"{type(exc).__name__}: {exc}"}

        report["sources"].append(info)
        if "error" in info:
            print(f"  ✗ {path.name}: {info['error']}")
        else:
            rows = info.get("row_count", "?")
            enc = info.get("encoding", "-")
            grains = ",".join(g["grain"] + "/" + g["confidence"] for g in info.get("grain_candidates", [])) or "-"
            print(f"  ✓ {path.name}  rows={rows} encoding={enc} grain={grains}")

    # 跨源汇总：同一语义角色出现在哪些文件的哪些列
    role_map: dict[str, list[dict[str, str]]] = {}
    for info in report["sources"]:
        for col in info.get("columns", []) or []:
            for role in col.get("role_candidates") or []:
                role_map.setdefault(role, []).append({"file": Path(info["path"]).name, "column": col["name"]})
    report["role_candidates"] = role_map

    # 同一角色出现多种列名 → 提示可能存在多候选口径
    report["ambiguity_flags"] = [
        {"role": role, "distinct_columns": len({h["column"] for h in hits}), "hits": hits}
        for role, hits in role_map.items()
        if len({h["column"] for h in hits}) > 1
    ]

    if args.target:
        tp = Path(args.target)
        try:
            report["target"] = probe_target_workbook(tp) if tp.suffix.lower() in {".xlsx", ".xlsm"} else probe_csv(tp)
            print(f"目标报表：{tp.name}")
        except Exception as exc:  # noqa: BLE001
            report["target"] = {"path": str(tp), "error": f"{type(exc).__name__}: {exc}"}
            print(f"  ✗ 目标报表探测失败：{exc}")

    if report.get("ambiguity_flags"):
        print("\n需要人工确认的口径歧义（同一语义对应多个列名）：")
        for flag in report["ambiguity_flags"]:
            cols = ", ".join(sorted({h["column"] for h in flag["hits"]}))
            print(f"  - {flag['role']}: {cols}")

    if args.out:
        out = Path(args.out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\n观测报告已写入 {out}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
